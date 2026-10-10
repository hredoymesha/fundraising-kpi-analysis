"""Excel report: Summary, By Region, By Dimension, Locations, Trends (with charts), Data Issues, Methods.

Pure presentation - every number comes from the engine's Result; nothing is recalculated here.
"""
from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .. import __version__
from .charts import add_trend_charts
from ..config import Config
from ..engine import Result

SHEETS = ["Summary", "By Region", "By Dimension", "Locations", "Trends", "Data Issues", "Methods"]

HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=12, color="1F3A5F")
MUTED = Font(color="666666", italic=True)
GREEN, AMBER, RED, GREY = (("C6EFCE", "006100"), ("FFEB9C", "7F6000"), ("FFC7CE", "9C0006"), ("E7E6E6", "404040"))
SCALE_RED, SCALE_AMBER, SCALE_GREEN = "F8696B", "FFEB84", "63BE7B"

# (key, header, number format) - columns absent from the data are skipped
REGION_COLUMNS = [
    ("region", "Region", None), ("shifts", "Shifts", "#,##0"), ("doors_knocked", "Doors knocked", "#,##0"),
    ("conversations", "Conversations", "#,##0"), ("pledges", "Pledges", "#,##0"),
    ("active_hours", "Active hours", "#,##0.0"), ("total_monthly_value_eur", "Monthly value (EUR)", "#,##0"),
    ("conversation_rate", "Conversation rate", "0.0%"), ("pledge_rate", "Pledge rate", "0.0%"),
    ("pledges_per_hour", "Pledges per hour", "0.00"), ("avg_monthly_donation_eur", "Avg monthly donation (EUR)", "#,##0.00"),
]
DIMENSION_COLUMNS = [("dimension", "Breakdown", None), ("value", "Value", None)] + REGION_COLUMNS[1:]
DIMENSION_LABELS = {"activity_type": "Activity type", "fundraiser_id": "Fundraiser", "location_id": "Location"}
CHANGE_FORMAT = "+0.0%;-0.0%;0.0%"
TREND_COLUMNS = [
    ("period", "Period", None), ("period_start", "Starts", "yyyy-mm-dd"), ("shifts", "Shifts", "#,##0"),
    ("pledges", "Pledges", "#,##0"), ("active_hours", "Active hours", "#,##0.0"),
    ("conversation_rate", "Conversation rate", "0.0%"), ("pledge_rate", "Pledge rate", "0.0%"),
    ("pledges_per_hour", "Pledges per hour", "0.00"), ("pph_rolling_4", "PPH, last 4 periods", "0.00"),
    ("avg_monthly_donation_eur", "Avg monthly donation (EUR)", "#,##0.00"),
    ("pledges_change_pct", "Pledges vs previous", CHANGE_FORMAT), ("pph_change_pct", "PPH vs previous", CHANGE_FORMAT),
]
LOCATION_COLUMNS = [
    ("location_id", "Location ID", None), ("location_label", "Location", None), ("location_type", "Type", None),
    ("pledges", "Pledges", "#,##0"), ("hours", "Hours", "#,##0.0"), ("target_pph", "Target PPH", "0.00"),
    ("achieved_pph", "Achieved PPH", "0.00"), ("target_attainment", "Target attainment", "0.0%"),
    ("productivity_gap", "PPH gap vs target", "+0.00;-0.00;0.00"),
    ("avg_monthly_donation_eur", "Avg monthly donation (EUR)", "#,##0.00"),
    ("planning_priority", "Planning tier", None), ("classification_note", "Why", None),
]
ISSUE_COLUMNS = [
    ("dataset", "Dataset", None), ("row", "Source row", "0"), ("column", "Column", None), ("rule", "Rule", None),
    ("severity", "Severity", None), ("value", "Value found", None), ("message", "What is wrong", None),
]
KPI_DEFINITIONS = [
    ("Conversation rate", "conversations ÷ doors knocked"),
    ("Pledge rate", "pledges ÷ conversations"),
    ("Pledges per hour (PPH)", "pledges ÷ active hours"),
    ("Target attainment", "achieved PPH ÷ target PPH (100% = exactly on target)"),
    ("Average monthly donation", "total monthly value ÷ pledges"),
]


def _py(v):
    """numpy / pandas value -> something openpyxl can write; missing -> empty cell."""
    if v is None or (not isinstance(v, (str, list, dict)) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, np.generic):
        return v.item()
    return v


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def _write_table(ws: Worksheet, df: pd.DataFrame, columns, *, freeze_col: bool = False) -> dict[str, str]:
    """Header in row 1, data below; styled header, frozen header row, filter, readable widths.
    Returns {key: column letter} for the columns written."""
    cols = [(k, h, f) for k, h, f in columns if k in df.columns]
    letters: dict[str, str] = {}
    for c, (key, header, _) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        letters[key] = cell.column_letter
    for r, row in enumerate(df[[k for k, _, _ in cols]].itertuples(index=False), start=2):
        for c, ((key, _, fmt), value) in enumerate(zip(cols, row), start=1):
            cell = ws.cell(row=r, column=c, value=_py(value))
            if cell.data_type == "f":
                # text from the user's file that starts with "=" would run as a formula when the report is opened
                cell.data_type = "s"
            if fmt:
                cell.number_format = fmt
    for c, (key, header, fmt) in enumerate(cols, start=1):
        longest = max([len(header) // 2 + 2] + [len(_shown(v, fmt)) for v in df[key]])
        ws.column_dimensions[letters[key]].width = max(11, min(longest + 2, 60))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2" if freeze_col else "A2"
    if len(df):
        ws.auto_filter.ref = f"A1:{ws.cell(row=len(df) + 1, column=len(cols)).coordinate}"
    return letters


def _shown(v, fmt: str | None) -> str:
    v = _py(v)
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v * 100:.1f}%" if fmt and "%" in fmt else f"{v:,.2f}"
    if isinstance(v, dt.datetime):
        return "2024-01-01"
    return str(v)


def _span(letter: str, n_rows: int) -> str:
    return f"{letter}2:{letter}{max(n_rows + 1, 2)}"


def _label_colors(config: Config) -> dict[str, tuple[str, str]]:
    """Best tier green, worst red, those between amber; 'insufficient' grey. Per-type tiers use the same scheme."""
    colors: dict[str, tuple[str, str]] = {config.classification.insufficient_label: GREY}
    for tiers in [config.classification.tiers, *config.classification.tiers_by_type.values()]:
        for i, tier in enumerate(tiers):
            colors.setdefault(tier.label, GREEN if i == 0 else RED if i == len(tiers) - 1 else AMBER)
    return colors


def _quote(text: str) -> str:
    return '"' + text.replace('"', '""') + '"'


def _scale(ws: Worksheet, rng: str, **kw) -> None:
    ws.conditional_formatting.add(rng, ColorScaleRule(start_color=SCALE_RED, mid_color=SCALE_AMBER,
                                                      end_color=SCALE_GREEN, **kw))


# ---- sheets ---------------------------------------------------------------------------------------

def _region_sheet(ws: Worksheet, result: Result) -> None:
    n = len(result.by_region)
    letters = _write_table(ws, result.by_region, REGION_COLUMNS, freeze_col=True)
    for key in ("pledges_per_hour", "conversation_rate", "pledge_rate"):
        if key in letters and n > 1:
            _scale(ws, _span(letters[key], n), start_type="min", mid_type="percentile", mid_value=50, end_type="max")
    if "pledges" in letters and n:
        ws.conditional_formatting.add(_span(letters["pledges"], n), DataBarRule(
            start_type="num", start_value=0, end_type="max", color="5B9BD5", showValue=True))


def _dimension_sheet(ws: Worksheet, result: Result) -> None:
    table = result.by_dimension
    if table is None or table.empty:
        ws["A1"] = "No breakdown columns in the data: map activity_type or fundraiser_id to see KPIs by those."
        ws["A1"].font = MUTED
        ws.column_dimensions["A"].width = 90
        return
    shown = table.assign(dimension=table["dimension"].map(lambda d: DIMENSION_LABELS.get(d, d)))
    letters = _write_table(ws, shown, DIMENSION_COLUMNS)
    # each breakdown gets its own colour scale: fundraisers should not be shaded against activity types
    start = 2
    for _, block in shown.groupby("dimension", sort=False):
        end = start + len(block) - 1
        if len(block) > 1:
            for key in ("pledges_per_hour", "conversation_rate", "pledge_rate"):
                if key in letters:
                    _scale(ws, f"{letters[key]}{start}:{letters[key]}{end}", start_type="min",
                           mid_type="percentile", mid_value=50, end_type="max")
        start = end + 1


def _trends_sheet(ws: Worksheet, result: Result) -> None:
    table = result.trends
    if table is None or table.empty:
        ws["A1"] = "No trends: the data needs a date column with readable dates."
        ws["A1"].font = MUTED
        ws.column_dimensions["A"].width = 90
        return
    n = len(table)
    letters = _write_table(ws, table, TREND_COLUMNS, freeze_col=True)
    for key in ("pledges_change_pct", "pph_change_pct"):
        if key in letters:
            rng = _span(letters[key], n)
            ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"],
                                                          font=Font(color=GREEN[1], bold=True)))
            ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                          font=Font(color=RED[1], bold=True)))
    cols = {k: ws[f"{letter}1"].column for k, letter in letters.items()}
    add_trend_charts(ws, cols, n, first_row=n + 4)


def _location_table(result: Result) -> pd.DataFrame | None:
    return result.by_location if result.by_location is not None else result.by_location_weekly


def _location_sheet(ws: Worksheet, result: Result, config: Config) -> None:
    table = _location_table(result)
    if table is None or table.empty:
        ws["A1"] = "No location data: add a 'locations' input (with target_pph) or a 'location_weekly' input."
        ws["A1"].font = MUTED
        ws.column_dimensions["A"].width = 90
        return
    n = len(table)
    letters = _write_table(ws, table, LOCATION_COLUMNS, freeze_col=True)
    positive = sorted({t.min for t in config.classification.tiers if t.min > 0})
    if "target_attainment" in letters and positive:
        mid = positive[0]
        top = positive[-1] if len(positive) > 1 else mid * 1.25
        _scale(ws, _span(letters["target_attainment"], n), start_type="num", start_value=0,
               mid_type="num", mid_value=mid, end_type="num", end_value=top)
    if "planning_priority" in letters:
        rng = _span(letters["planning_priority"], n)
        for label, (bg, fg) in _label_colors(config).items():
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[_quote(label)],
                                                          fill=_fill(bg), font=Font(color=fg, bold=True)))
    if "classification_note" in letters:
        ws.column_dimensions[letters["classification_note"]].width = 58


def _issues_sheet(ws: Worksheet, result: Result) -> None:
    issues = result.issues
    letters = _write_table(ws, issues, ISSUE_COLUMNS)
    if issues.empty:
        ws["A2"] = "No data issues found."
        ws["A2"].font = MUTED
        return
    ws.column_dimensions[letters["message"]].width = 70
    rng = _span(letters["severity"], len(issues))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"error"'], fill=_fill(RED[0]),
                                                  font=Font(color=RED[1], bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"warning"'], fill=_fill(AMBER[0]),
                                                  font=Font(color=AMBER[1], bold=True)))


def _section(ws: Worksheet, row: int, title: str) -> int:
    ws.cell(row=row, column=1, value=title).font = SECTION_FONT
    return row + 1


def _summary_sheet(ws: Worksheet, result: Result, config: Config) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "F2F Fundraising KPI Report"
    ws["A1"].font = Font(bold=True, size=16, color="1F3A5F")
    ws["A2"] = f"Generated {dt.datetime.now():%Y-%m-%d %H:%M} · f2f-kpi {__version__}"
    ws["A2"].font = MUTED

    row = _section(ws, 4, "Headline KPIs")
    o = result.overall.iloc[0] if len(result.overall) else pd.Series(dtype=float)
    headline = [
        ("Pledges", "pledges", "#,##0"), ("Active hours", "active_hours", "#,##0.0"),
        ("Doors knocked", "doors_knocked", "#,##0"), ("Conversations", "conversations", "#,##0"),
        ("Conversation rate", "conversation_rate", "0.0%"), ("Pledge rate", "pledge_rate", "0.0%"),
        ("Pledges per hour", "pledges_per_hour", "0.00"),
        ("Total monthly value (EUR)", "total_monthly_value_eur", "#,##0"),
        ("Avg monthly donation (EUR)", "avg_monthly_donation_eur", "#,##0.00"),
    ]
    for label, key, fmt in headline:
        ws.cell(row=row, column=1, value=label)
        value = _py(o.get(key)) if key in o.index else None
        cell = ws.cell(row=row, column=2, value=value)
        cell.number_format = fmt
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right")
        if value is None:
            ws.cell(row=row, column=3, value=result.skipped_kpis.get(key, "not available (no data)")).font = MUTED
        row += 1

    row = _section(ws, row + 1, "Key findings")
    for line in _findings(result, config):
        ws.cell(row=row, column=1, value="• " + line)
        row += 1

    row = _section(ws, row + 1, "Data quality")
    for c, h in enumerate(["Dataset", "Rows read", "Rows used"], start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    row += 1
    for name, read in result.rows_read.items():
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=read)
        ws.cell(row=row, column=3, value=result.rows_kept.get(name)).alignment = Alignment(horizontal="left")
        row += 1
    errors = int((result.issues["severity"] == "error").sum())
    warnings = int((result.issues["severity"] == "warning").sum())
    for label, count, colors in [("Errors (rows excluded)", errors, RED), ("Warnings (rows kept)", warnings, AMBER)]:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=count)
        cell.font = Font(bold=True)
        ws.conditional_formatting.add(cell.coordinate, CellIsRule(operator="greaterThan", formula=["0"],
                                                                  fill=_fill(colors[0]), font=Font(color=colors[1], bold=True)))
        row += 1
    if not result.issues.empty:
        row += 1
        ws.cell(row=row, column=1, value="Issues by rule (see the Data Issues sheet)").font = Font(bold=True)
        row += 1
        for rule, count in result.issues["rule"].value_counts().items():
            ws.cell(row=row, column=1, value=rule)
            ws.cell(row=row, column=2, value=int(count))
            row += 1


def _findings(result: Result, config: Config) -> list[str]:
    lines: list[str] = []
    r = result.by_region
    if len(r) and "pledges" in r:
        top = r.iloc[0]
        lines.append(f"Most pledges: {top['region']} with {int(top['pledges']):,}.")
    if len(r) and "pledges_per_hour" in r and r["pledges_per_hour"].notna().any():
        best = r.loc[r["pledges_per_hour"].idxmax()]
        lines.append(f"Most productive region: {best['region']} at {best['pledges_per_hour']:.2f} pledges per hour.")
    loc = _location_table(result)
    if loc is not None and len(loc):
        scored = loc[loc["target_attainment"].notna()]
        if len(scored):
            best = scored.iloc[0]
            name = best.get("location_label") if "location_label" in scored else None
            lines.append(f"Best location vs target: {name if isinstance(name, str) else best['location_id']} "
                         f"at {best['target_attainment']:.0%} of target PPH.")
        counts = loc["planning_priority"].value_counts()
        tiers = [t.label for t in config.classification.tiers] + [config.classification.insufficient_label]
        parts = [f"{counts[t]} {t}" for t in tiers if t in counts]
        lines.append("Locations: " + ", ".join(parts) + ".")
        worst = config.classification.tiers[-1].label
        review = loc[loc["planning_priority"] == worst]
        if len(review):
            names = [str(x) for x in (review["location_label"] if "location_label" in review else review["location_id"])]
            more = f" and {len(names) - 5} more" if len(names) > 5 else ""
            lines.append(f"To review: {', '.join(names[:5])}{more}.")
    return lines or ["No findings: no usable rows."]


def _methods_sheet(ws: Worksheet, result: Result, config: Config) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    row = _section(ws, 1, "KPI definitions")
    for name, text in KPI_DEFINITIONS:
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=text)
        row += 1
    notes = [
        "Group KPIs divide summed totals (e.g. all pledges ÷ all hours), never average row-level rates.",
        "A KPI with a zero or missing denominator is left blank, not shown as 0.",
        "Location target PPH is the hours-weighted mean of the target PPH of its rows.",
        "Trends use calendar weeks (Monday start) or months, including quiet periods. 'Last 4 periods' is "
        "pledges ÷ hours over the four most recent periods; changes compare with the period just before.",
    ]
    for text in notes:
        ws.cell(row=row, column=1, value="Note")
        ws.cell(row=row, column=2, value=text)
        row += 1

    v, c = config.validation, config.classification
    row = _section(ws, row + 1, "Validation settings")
    rng = f"{v.expected_date_range[0]} to {v.expected_date_range[1]}" if v.expected_date_range else \
        f"within {v.date_window_days} days of the median date"
    for k, text in [("Max hours per row", f"{v.max_hours_per_row:g}"), ("Max monthly value per pledge (EUR)", f"{v.max_value_per_pledge:g}"),
                    ("Expected dates", rng), ("achieved_pph tolerance", f"{v.achieved_pph_tolerance:g}"),
                    ("Severity overrides", ", ".join(f"{r}={s}" for r, s in v.severity.items()) or "none")]:
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=text)
        row += 1

    row = _section(ws, row + 1, "Location classification")
    ws.cell(row=row, column=1, value="Minimum hours")
    ws.cell(row=row, column=2, value=f"{c.min_hours:g} (below this: '{c.insufficient_label}')")
    row += 1
    for t in c.tiers:
        ws.cell(row=row, column=1, value=t.label)
        ws.cell(row=row, column=2, value=f"target attainment of {t.min:.0%} or more")
        row += 1
    for ltype, tiers in c.tiers_by_type.items():
        for t in tiers:
            ws.cell(row=row, column=1, value=f"{ltype}: {t.label}")
            ws.cell(row=row, column=2, value=f"{t.min:.0%} or more")
            row += 1

    row = _section(ws, row + 1, "Input files")
    for name, spec in config.inputs.items():
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=str(spec.path))
        row += 1


def _build(result: Result, config: Config) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in SHEETS}
    _summary_sheet(sheets["Summary"], result, config)
    _region_sheet(sheets["By Region"], result)
    _dimension_sheet(sheets["By Dimension"], result)
    _location_sheet(sheets["Locations"], result, config)
    _trends_sheet(sheets["Trends"], result)
    _issues_sheet(sheets["Data Issues"], result)
    _methods_sheet(sheets["Methods"], result, config)
    return wb


def report_bytes(result: Result, config: Config) -> bytes:
    """The finished workbook as bytes, built in memory so nothing touches the disk."""
    buffer = io.BytesIO()
    _build(result, config).save(buffer)
    return buffer.getvalue()


def write_report(result: Result, config: Config, path: Path | str) -> Path:
    """Write the workbook and return its path."""
    wb = _build(result, config)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
