"""Read CSV / Excel into an all-text DataFrame with source line numbers. No interpretation happens here."""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path
from typing import BinaryIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .config import InputSpec

ROW = "_row"  # line number in the source file (what a person sees in Excel / a text editor)
EXCEL_SUFFIXES = (".xlsx", ".xlsm")
TEXT_SUFFIXES = (".csv", ".txt", ".tsv")


class FileReadError(ValueError):
    """A file could not be read as a table. The message never includes cell contents, so it is safe to show or log."""


def _excel_cell_text(value: object, decimal: str) -> str:
    """Real Excel numbers/dates arrive as Python objects; render them the way the numeric/date parsing expects.
    Numbers use the configured decimal character (so a 4.5 cell is not misread when decimal is ',')."""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        text = repr(value)
        return text.replace(".", decimal)
    return str(value)  # dates -> '2024-01-09 00:00:00'


def _parse(source: Path | BinaryIO, suffix: str, label: str, *, sheet: str | int | None, header_row: int,
           delimiter: str | None, encoding: str, decimal: str) -> pd.DataFrame:
    skip = header_row - 1
    try:
        if suffix in EXCEL_SUFFIXES:
            df = pd.read_excel(source, sheet_name=sheet if sheet is not None else 0, header=skip, dtype=object,
                               keep_default_na=False, engine="openpyxl")
            df = df.apply(lambda col: col.map(lambda v: _excel_cell_text(v, decimal)))
        elif suffix in TEXT_SUFFIXES:
            df = pd.read_csv(source, header=skip, dtype=str, keep_default_na=False, skip_blank_lines=False,
                             sep=delimiter or ("\t" if suffix == ".tsv" else None), engine="python", encoding=encoding)
        else:
            raise FileReadError(f"unsupported file type '{suffix or '(none)'}' for {label} (use .csv or .xlsx)")
    except FileReadError:
        raise
    except pd.errors.EmptyDataError:
        raise FileReadError(f"{label} is empty.") from None
    except UnicodeDecodeError:
        raise FileReadError(f"{label} is not readable as {encoding.replace('-sig', '')} text. Try another encoding, "
                            f"or save it again as CSV (UTF-8).") from None
    except (pd.errors.ParserError, csv.Error):  # csv.Error: no separator could be detected
        raise FileReadError(f"{label} could not be split into columns. Check the separator and the header row.") from None
    except (zipfile.BadZipFile, InvalidFileException):
        raise FileReadError(f"{label} is not a valid Excel workbook (.xlsx). Is it damaged, or a different format?") from None
    except ValueError as exc:  # pandas raises a plain ValueError for an unknown sheet name
        if "Worksheet" in str(exc) or "sheet" in str(exc).lower():
            raise FileReadError(f"{label} has no sheet called '{sheet}'.") from None
        raise FileReadError(f"{label} could not be read as a table.") from None
    df = df.fillna("").astype(str)  # the python CSV engine yields NaN for blank lines / short rows
    df.columns = [str(c).strip() for c in df.columns]
    df[ROW] = range(header_row + 1, header_row + 1 + len(df))
    if len(df.columns) == 1:  # only the row-number column: no headers at all
        raise FileReadError(f"{label} has no column headings on row {header_row}.")
    blank = df.drop(columns=[ROW]).apply(lambda col: col.str.strip() == "").all(axis=1)
    return df.loc[~blank].reset_index(drop=True)


def read_table(spec: InputSpec, decimal: str = ".") -> pd.DataFrame:
    if not spec.path.exists():
        raise FileNotFoundError(f"input file not found: {spec.path}")
    return _parse(spec.path, spec.path.suffix.lower(), str(spec.path), sheet=spec.sheet, header_row=spec.header_row,
                  delimiter=spec.delimiter, encoding=spec.encoding, decimal=decimal)


def read_bytes(data: bytes, filename: str, *, sheet: str | int | None = None, header_row: int = 1,
               delimiter: str | None = None, encoding: str = "utf-8-sig", decimal: str = ".") -> pd.DataFrame:
    """Same as read_table, for a file that is already in memory (e.g. an upload). Nothing is written to disk."""
    if not data.strip():
        raise FileReadError(f"'{filename}' is empty.")
    return _parse(io.BytesIO(data), Path(filename).suffix.lower(), f"'{filename}'", sheet=sheet,
                  header_row=header_row, delimiter=delimiter, encoding=encoding, decimal=decimal)


def excel_sheet_names(data: bytes, filename: str = "the workbook") -> list[str]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError):
        raise FileReadError(f"'{filename}' is not a valid Excel workbook (.xlsx).") from None
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()
