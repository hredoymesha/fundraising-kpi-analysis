"""Filtering, reading from memory, and the in-memory report: the engine pieces the web app relies on.

Hand-calculated data (blank date on the last row on purpose):
  2024-01-08  N  L1  doors 100  conv 30  pledges 6   value 90   hours 4
  2024-01-09  N  L1  doors 300  conv 60  pledges 12  value 180  hours 8
  2024-01-15  S  L2  doors 200  conv 40  pledges 4   value 80   hours 5
  (no date)   S  L2  doors 50   conv 10  pledges 2   value 30   hours 2
  Targets: L1 = 1.0 PPH, L2 = 2.0 PPH.
"""
import io
import math

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from f2f_kpi import analyse, config_from_dict, filter_options, report_bytes, with_filters, write_report
from f2f_kpi.engine import filter_activity
from f2f_kpi.io import FileReadError, excel_sheet_names, read_bytes, read_table
from f2f_kpi.config import InputSpec

ACTIVITY = (
    "date,region,location_id,doors_knocked,conversations,pledges,total_monthly_value_eur,active_hours\n"
    "2024-01-08,N,L1,100,30,6,90,4\n"
    "2024-01-09,N,L1,300,60,12,180,8\n"
    "2024-01-15,S,L2,200,40,4,80,5\n"
    ",S,L2,50,10,2,30,2\n"
)
LOCATIONS = "location_id,location_label,location_type,target_pph\nL1,Alpha,Stand,1.0\nL2,Beta,Fair,2.0\n"


@pytest.fixture
def setup():
    config = config_from_dict({}, require_activity=False)
    raw = {"activity": read_bytes(ACTIVITY.encode(), "a.csv"), "locations": read_bytes(LOCATIONS.encode(), "l.csv")}
    return config, analyse(config, raw)


# ---- filtering --------------------------------------------------------------------------------------

def test_full_data_numbers_by_hand(setup):
    _, r = setup
    o = r.overall.iloc[0]
    assert o["pledges"] == 24 and o["active_hours"] == 19  # 6+12+4+2 ; 4+8+5+2
    assert o["pledges_per_hour"] == pytest.approx(24 / 19)
    # attainment: pledges 24 over expected 1.0*12 + 2.0*7 = 26
    assert r.location_overall.iloc[0]["target_attainment"] == pytest.approx(24 / 26)


def test_region_filter_recalculates_everything(setup):
    config, r = setup
    f = with_filters(r, config, regions=["N"])
    o = f.overall.iloc[0]
    # N: doors 400, conversations 90, pledges 18, value 270, hours 12
    assert (o["conversation_rate"], o["pledge_rate"], o["pledges_per_hour"], o["avg_monthly_donation_eur"]) == \
        pytest.approx((90 / 400, 18 / 90, 18 / 12, 270 / 18))
    assert list(f.by_region["region"]) == ["N"]
    assert list(f.by_location["location_id"]) == ["L1"]
    assert f.location_overall.iloc[0]["target_attainment"] == pytest.approx(18 / 12)  # 1.5
    s = with_filters(r, config, regions=["S"])
    # S: target hours-weighted (2*5 + 2*2) / 7 = 2.0 ; achieved 6 / 7
    assert s.location_overall.iloc[0]["target_attainment"] == pytest.approx((6 / 7) / 2.0)


def test_date_filter_includes_both_ends_and_drops_undated_rows(setup):
    config, r = setup
    f = with_filters(r, config, start="2024-01-09", end="2024-01-15")
    # only the 9th and the 15th: pledges 12 + 4, hours 8 + 5 ; the undated row is out
    assert f.overall.iloc[0]["pledges"] == 16 and f.overall.iloc[0]["active_hours"] == 13
    assert with_filters(r, config, start="2024-01-09").overall.iloc[0]["pledges"] == 16  # open-ended


def test_no_filter_keeps_every_row_including_undated(setup):
    config, r = setup
    f = with_filters(r, config)
    assert len(f.activity) == 4 and f.overall.iloc[0]["pledges"] == 24


def test_trends_follow_the_filter(setup):
    config, r = setup
    f = with_filters(r, config, regions=["S"])
    assert list(f.trends["period"]) == ["2024-W03"]  # the dated S row only; the undated one has no week
    assert f.trends.iloc[0]["pledges"] == 4


def test_filter_that_matches_nothing_gives_empty_tables_not_an_error(setup):
    config, r = setup
    f = with_filters(r, config, regions=[])
    assert f.activity.empty and f.overall.empty and f.by_region.empty
    assert f.by_location is None and f.location_overall is None and f.trends is None
    assert len(f.issues) == len(r.issues) and f.rows_read == r.rows_read  # the upload's story is unchanged


def test_filtering_does_not_change_the_original_result(setup):
    config, r = setup
    with_filters(r, config, regions=["N"], start="2024-01-09")
    assert len(r.activity) == 4 and r.overall.iloc[0]["pledges"] == 24


def test_filter_options(setup):
    _, r = setup
    opts = filter_options(r)
    assert opts["regions"] == ["N", "S"]
    assert opts["date_min"] == pd.Timestamp("2024-01-08") and opts["date_max"] == pd.Timestamp("2024-01-15")


def test_filter_options_without_dates_or_regions():
    config = config_from_dict({}, require_activity=False)
    r = analyse(config, {"activity": read_bytes(b"region,pledges,active_hours\nA,1,1\n", "a.csv")})
    assert filter_options(r) == {"regions": ["A"], "date_min": None, "date_max": None}
    assert filter_activity(r.activity, None, "2024-01-01", None).equals(r.activity)  # no date column: nothing to filter


# ---- reading from memory ----------------------------------------------------------------------------

def test_read_bytes_matches_read_table(tmp_path):
    (tmp_path / "a.csv").write_text(ACTIVITY, encoding="utf-8")
    from_disk = read_table(InputSpec(path=tmp_path / "a.csv"))
    from_memory = read_bytes(ACTIVITY.encode(), "a.csv")
    pd.testing.assert_frame_equal(from_disk, from_memory)


@pytest.mark.parametrize("data,name,expected", [
    (b"", "empty.csv", "is empty"),
    (b"   \n  ", "blank.csv", "is empty"),
    (bytes([0xff, 0xfe, 0x00, 0x81]) * 50, "binary.csv", "not readable as utf-8"),
    (b"this is not a zip", "broken.xlsx", "not a valid Excel workbook"),
    (b"%PDF-1.4", "scan.pdf", "unsupported file type '.pdf'"),
    (b"data", "noextension", "unsupported file type"),
])
def test_unreadable_files_get_a_plain_message(data, name, expected):
    with pytest.raises(FileReadError, match=expected) as exc:
        read_bytes(data, name)
    assert name in str(exc.value)


def test_error_messages_never_contain_file_contents():
    secret = b"Gebiet;Zusagen\nSecretTownName;12345\n"
    for kwargs in ({"encoding": "ascii"}, {"sheet": "x"}):
        try:
            read_bytes(secret + bytes([0xff]), "f.csv", **kwargs)
        except FileReadError as exc:
            assert b"SecretTownName" not in str(exc).encode() and "12345" not in str(exc)


def test_a_separator_that_cannot_be_found_gets_a_plain_message(monkeypatch):
    import csv as csv_module

    def cannot_sniff(*args, **kwargs):
        raise csv_module.Error("Could not determine delimiter")

    monkeypatch.setattr(pd, "read_csv", cannot_sniff)
    with pytest.raises(FileReadError, match="could not be split into columns"):
        read_bytes(b"whatever", "odd.csv")


def test_header_only_file_reads_as_zero_rows():
    df = read_bytes(b"region,pledges,active_hours\n", "h.csv")
    assert len(df) == 0 and {"region", "pledges", "active_hours"} <= set(df.columns)


def _workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    ws.append(["nothing here"])
    daily = wb.create_sheet("Daily")
    daily.append(["Export"])
    daily.append(["region", "pledges", "active_hours"])
    daily.append(["A", 10, 4.5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_sheets_can_be_listed_and_chosen_from_memory():
    data = _workbook()
    assert excel_sheet_names(data) == ["Cover", "Daily"]
    df = read_bytes(data, "b.xlsx", sheet="Daily", header_row=2)
    assert list(df["pledges"]) == ["10"] and list(df["_row"]) == [3]
    with pytest.raises(FileReadError, match="no sheet called 'Missing'"):
        read_bytes(data, "b.xlsx", sheet="Missing", header_row=2)


def test_excel_numbers_follow_the_decimal_mark_from_memory():
    assert list(read_bytes(_workbook(), "b.xlsx", sheet="Daily", header_row=2, decimal=",")["active_hours"]) == ["4,5"]


def test_excel_sheet_names_rejects_a_damaged_workbook():
    with pytest.raises(FileReadError, match="not a valid Excel workbook"):
        excel_sheet_names(b"nope", "x.xlsx")


def test_analyse_needs_activity_data():
    from f2f_kpi import ConfigError
    with pytest.raises(ConfigError, match="no activity data"):
        analyse(config_from_dict({}, require_activity=False), {})


# ---- report from memory -----------------------------------------------------------------------------

def test_report_bytes_is_the_same_workbook_write_report_makes(setup, tmp_path):
    config, r = setup
    data = report_bytes(r, config)
    from_bytes = load_workbook(io.BytesIO(data))
    from_file = load_workbook(write_report(r, config, tmp_path / "r.xlsx"))
    assert from_bytes.sheetnames == from_file.sheetnames == [
        "Summary", "By Region", "By Dimension", "Locations", "Trends", "Data Issues", "Methods"]
    for name in ("By Region", "Locations", "Trends"):
        assert [c.value for row in from_bytes[name].iter_rows() for c in row] == \
               [c.value for row in from_file[name].iter_rows() for c in row]
    assert list(tmp_path.iterdir()) == [tmp_path / "r.xlsx"]  # report_bytes itself wrote nothing


def test_report_bytes_numbers_by_hand(setup):
    config, r = setup
    ws = load_workbook(io.BytesIO(report_bytes(r, config)))["Summary"]
    s = {row[0]: row[1] for row in ws.iter_rows(max_col=2, values_only=True) if row[0]}
    assert s["Pledges"] == 24 and s["Pledges per hour"] == pytest.approx(24 / 19)
    assert s["Avg monthly donation (EUR)"] == pytest.approx(380 / 24)
    assert math.isclose(s["Conversation rate"], 140 / 650)
