"""Loading CSV/XLSX, config-based column mapping, and config errors."""
import pandas as pd
import pytest
from openpyxl import Workbook

from f2f_kpi import ConfigError, config_from_dict, load_config, run
from f2f_kpi.io import read_table
from f2f_kpi.config import InputSpec


def test_csv_with_other_headers_semicolons_decimal_comma_and_german_dates(tmp_path):
    (tmp_path / "export.csv").write_text(
        "Datum;Gebiet;Türen;Gespräche;Zusagen;Euro;Stunden\n"
        "09.01.2024;Nord;100;30;6;90,50;4,5\n"
        "10.01.2024;Nord;300;60;12;180;8\n"
        "10.01.2024;Süd;200;40;4;80;5\n", encoding="utf-8")
    cfg = config_from_dict({
        "inputs": {"activity": {"path": "export.csv", "delimiter": ";"}},
        "columns": {"activity": {"date": "Datum", "region": ["Region", "Gebiet"], "doors_knocked": "Türen",
                                 "conversations": "Gespräche", "pledges": "Zusagen",
                                 "total_monthly_value_eur": "Euro", "active_hours": "Stunden"}},
        "formats": {"date": "%d.%m.%Y", "decimal": ","},
    }, tmp_path)
    res = run(cfg)
    assert res.issues.empty
    nord = res.by_region.set_index("region").loc["Nord"]
    assert nord["active_hours"] == pytest.approx(12.5)  # 4.5 + 8
    assert nord["total_monthly_value_eur"] == pytest.approx(270.5)  # 90.50 + 180
    assert nord["pledges_per_hour"] == pytest.approx(18 / 12.5)  # 1.44


def test_header_matching_ignores_case_and_surrounding_spaces(tmp_path):
    (tmp_path / "a.csv").write_text(" REGION ,Pledges,ACTIVE_HOURS\nA,4,2\n", encoding="utf-8")
    res = run(config_from_dict({"inputs": {"activity": "a.csv"}}, tmp_path))
    assert res.issues.empty and res.overall.loc[0, "pledges_per_hour"] == pytest.approx(2.0)


def test_xlsx_sheet_and_header_row_and_source_row_numbers(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    ws2 = wb.create_sheet("Daily")
    ws2.append(["Weekly export - team X"])
    ws2.append([])
    ws2.append(["region", "pledges", "active_hours"])  # header on row 3
    ws2.append(["A", 10, 5])  # row 4
    ws2.append(["A", "oops", 5])  # row 5
    wb.save(tmp_path / "book.xlsx")
    spec = InputSpec(path=tmp_path / "book.xlsx", sheet="Daily", header_row=3)
    df = read_table(spec)
    assert list(df["_row"]) == [4, 5]  # Excel row numbers, so an issue points at the right cell
    res = run(config_from_dict({"inputs": {"activity": {"path": "book.xlsx", "sheet": "Daily", "header_row": 3}}}, tmp_path))
    bad = res.issues.iloc[0]
    assert (bad["rule"], bad["row"], bad["column"], bad["value"]) == ("non_numeric_value", 5, "pledges", "oops")
    assert res.overall.loc[0, "pledges"] == 10  # the bad row is excluded


def test_missing_required_column_is_reported_not_crashed(tmp_path):
    (tmp_path / "a.csv").write_text("region,pledges\nA,4\n", encoding="utf-8")  # no hours column
    res = run(config_from_dict({"inputs": {"activity": "a.csv"}}, tmp_path))
    miss = res.issues[res.issues["rule"] == "missing_column"]
    assert list(miss["column"]) == ["active_hours"] and miss.iloc[0]["severity"] == "error"
    assert res.activity.empty


def test_severity_can_be_downgraded_or_ignored(tmp_path):
    (tmp_path / "a.csv").write_text("region,pledges,active_hours,conversations,doors_knocked\n"
                                    "A,5,2,50,40\nA,5,2,10,40\n", encoding="utf-8")  # row 2: 50 > 40
    base = {"inputs": {"activity": "a.csv"}}
    assert run(config_from_dict(base, tmp_path)).activity["pledges"].sum() == 5  # error: row dropped
    warn = {**base, "validation": {"severity": {"conversations_gt_doors": "warning"}}}
    res = run(config_from_dict(warn, tmp_path))
    assert res.activity["pledges"].sum() == 10 and res.issues.iloc[0]["severity"] == "warning"
    ign = {**base, "validation": {"severity": {"conversations_gt_doors": "ignore"}}}
    assert run(config_from_dict(ign, tmp_path)).issues.empty


def test_blank_lines_do_not_shift_reported_row_numbers(tmp_path):
    (tmp_path / "a.csv").write_text("region,pledges,active_hours\nA,1,1\n\nA,x,1\n", encoding="utf-8")
    res = run(config_from_dict({"inputs": {"activity": "a.csv"}}, tmp_path))
    assert res.issues.iloc[0]["row"] == 4  # the bad value really is on line 4


def test_missing_input_file_raises_clearly(tmp_path):
    with pytest.raises(FileNotFoundError, match="nope.csv"):
        run(config_from_dict({"inputs": {"activity": "nope.csv"}}, tmp_path))


@pytest.mark.parametrize("raw,msg", [
    ({"inputs": {"activity": "a.csv"}, "colums": {}}, "unknown top-level"),
    ({"inputs": {}}, "inputs.activity is required"),
    ({"inputs": {"activity": "a.csv"}, "columns": {"activity": {"pledgez": "x"}}}, "not a known field"),
    ({"inputs": {"activity": "a.csv"}, "validation": {"severity": {"made_up": "error"}}}, "unknown rule"),
    ({"inputs": {"activity": "a.csv"}, "validation": {"severity": {"negative_hours": "fatal"}}}, "error, warning or ignore"),
    ({"inputs": {"activity": "a.csv"}, "formats": {"decimal": ";"}}, "decimal"),
])
def test_config_mistakes_give_clear_errors(raw, msg):
    with pytest.raises(ConfigError, match=msg):
        config_from_dict(raw)


def test_load_config_resolves_paths_relative_to_the_yaml_file(tmp_path):
    (tmp_path / "cfg").mkdir()
    (tmp_path / "cfg" / "c.yaml").write_text("inputs:\n  activity: ../data.csv\n", encoding="utf-8")
    cfg = load_config(tmp_path / "cfg" / "c.yaml")
    assert cfg.inputs["activity"].path == tmp_path / "cfg" / ".." / "data.csv"


def test_bad_yaml_is_a_config_error(tmp_path):
    (tmp_path / "c.yaml").write_text("inputs: [unclosed\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="not valid YAML"):
        load_config(tmp_path / "c.yaml")


# ---- regressions found in the pre-Step-4 review --------------------------------------------------

def _mixed_workbook(path):
    import datetime as dt
    wb = Workbook()
    ws = wb.active
    ws.append(["date", "region", "pledges", "active_hours"])
    ws.append([dt.datetime(2024, 1, 9), "A", 10, 4.5])  # real date and number cells
    ws.append(["10.01.2024", "A", "12", "8,5"])  # text cells in the team's own format
    wb.save(path)


def test_excel_number_cells_are_not_misread_when_decimal_is_comma(tmp_path):
    _mixed_workbook(tmp_path / "x.xlsx")
    res = run(config_from_dict({"inputs": {"activity": "x.xlsx"},
                                "formats": {"decimal": ",", "date": "%d.%m.%Y"}}, tmp_path))
    assert res.issues.empty
    assert sorted(res.activity["active_hours"]) == [4.5, 8.5]  # was 45.0 before the fix
    assert res.overall.loc[0, "pledges_per_hour"] == pytest.approx(22 / 13)


def test_excel_date_cells_survive_a_custom_date_format(tmp_path):
    _mixed_workbook(tmp_path / "x.xlsx")
    res = run(config_from_dict({"inputs": {"activity": "x.xlsx"},
                                "formats": {"decimal": ",", "date": "%d.%m.%Y"}}, tmp_path))
    assert list(res.activity["date"].dt.strftime("%Y-%m-%d")) == ["2024-01-09", "2024-01-10"]


def test_row_counts_read_and_kept_are_tracked(tmp_path):
    (tmp_path / "a.csv").write_text("region,pledges,active_hours\nA,1,1\nA,x,1\nA,2,1\n", encoding="utf-8")
    res = run(config_from_dict({"inputs": {"activity": "a.csv"}}, tmp_path))
    assert res.rows_read == {"activity": 3} and res.rows_kept == {"activity": 2}


# ---- found in the final review ---------------------------------------------------------------------

def test_dates_in_an_unsupported_style_are_reported_not_a_crash(tmp_path):
    (tmp_path / "a.csv").write_text("date,region,pledges,active_hours\n5/1/2024,A,1,1\n6/1/2024,A,2,1\n", encoding="utf-8")
    res = run(config_from_dict({"inputs": {"activity": "a.csv"}}, tmp_path))  # used to die on a missing 'median date'
    assert list(res.issues["rule"]) == ["unparseable_date", "unparseable_date"]
    assert "formats.date" in res.issues.iloc[0]["message"] and res.activity.empty


def test_missing_column_message_names_the_likely_match_and_hides_internals(tmp_path):
    (tmp_path / "a.csv").write_text("region,pledges,Active Hours\nA,1,1\n", encoding="utf-8")
    res = run(config_from_dict({"inputs": {"activity": "a.csv"}}, tmp_path))
    message = res.issues.iloc[0]["message"]
    assert "active_hours" in message and "'Active Hours' looks like a match" in message
    assert "columns.activity.active_hours" in message and "_row" not in message


def test_a_location_id_repeated_with_other_details_is_reported(tmp_path):
    (tmp_path / "a.csv").write_text("region,location_id,pledges,active_hours\nA,L1,6,3\n", encoding="utf-8")
    (tmp_path / "l.csv").write_text("location_id,target_pph\nL1,1.0\nL1,3.0\nL1,1.0\n", encoding="utf-8")
    res = run(config_from_dict({"inputs": {"activity": "a.csv", "locations": "l.csv"}}, tmp_path))
    rules = [(i.row, i.rule) for i in res.issues.itertuples()]
    assert rules == [(4, "exact_duplicate_row"), (3, "duplicate_location_id")]  # row 4 copies row 2 exactly; row 3 differs
    # the first row wins: 6 pledges / 3 hours = 2.0 against a target of 1.0
    assert res.by_location.iloc[0]["target_attainment"] == pytest.approx(2.0)
