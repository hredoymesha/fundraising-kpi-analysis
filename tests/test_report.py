"""Excel report tests. A tiny dataset with hand-calculated numbers, plus the full sample."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from f2f_kpi import config_from_dict, load_config, run, write_report
from f2f_kpi.report import SHEETS

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def small(tmp_path):
    """Hand-calculated inputs.
    A/L1: doors 100+300 (+10 in the bad row, dropped), conversations 30+60, pledges 6+12, value 90+180, hours 4+8
        -> doors 400, conv 90, pledges 18, value 270, hours 12 -> 0.225, 0.2, 1.5 PPH, 15.00 EUR ; L1 target 1.0 -> 150%
    B/L2: doors 200, conv 40, pledges 4, value 80, hours 5 -> 0.2, 0.1, 0.8 PPH, 20.00 EUR ; L2 target 1.0 -> 80%
    Row 5 has more conversations than doors -> one error issue, row excluded.
    """
    (tmp_path / "shifts.csv").write_text(
        "region,location_id,doors_knocked,conversations,pledges,total_monthly_value_eur,active_hours\n"
        "A,L1,100,30,6,90,4\n"
        "A,L1,300,60,12,180,8\n"
        "B,L2,200,40,4,80,5\n"
        "A,L1,10,20,1,15,1\n", encoding="utf-8")
    (tmp_path / "locations.csv").write_text(
        "location_id,location_label,location_type,target_pph\nL1,Alpha,Stand,1.0\nL2,Beta,Fair,1.0\n", encoding="utf-8")
    cfg = config_from_dict({"inputs": {"activity": "shifts.csv", "locations": "locations.csv"},
                            "classification": {"min_hours": 1}}, tmp_path)
    result = run(cfg)
    path = write_report(result, cfg, tmp_path / "out" / "report.xlsx")
    return load_workbook(path), result, cfg


def table(ws):
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def summary_values(ws):
    return {r[0]: r[1] for r in ws.iter_rows(max_col=2, values_only=True) if r[0]}


def test_sheets_exist_in_the_documented_order(small):
    wb, _, _ = small
    assert wb.sheetnames == SHEETS == ["Summary", "By Region", "By Dimension", "Locations", "Trends",
                                       "Data Issues", "Methods"]


def test_summary_headline_numbers_by_hand(small):
    s = summary_values(small[0]["Summary"])
    # totals over the 3 valid rows: doors 600, conversations 130, pledges 22, value 350, hours 17
    assert s["Pledges"] == 22 and s["Doors knocked"] == 600 and s["Conversations"] == 130
    assert s["Active hours"] == pytest.approx(17)
    assert s["Conversation rate"] == pytest.approx(130 / 600)
    assert s["Pledge rate"] == pytest.approx(22 / 130)
    assert s["Pledges per hour"] == pytest.approx(22 / 17)
    assert s["Avg monthly donation (EUR)"] == pytest.approx(350 / 22)


def test_summary_data_quality_block(small):
    s = summary_values(small[0]["Summary"])
    assert s["activity"] == 4  # rows read ...
    assert small[0]["Summary"]["C" + str(next(c.row for c in small[0]["Summary"]["A"] if c.value == "activity"))].value == 3  # ... and used
    assert s["Errors (rows excluded)"] == 1 and s["Warnings (rows kept)"] == 0
    assert s["conversations_gt_doors"] == 1


def test_summary_findings_name_the_top_region_and_locations_to_review(small):
    text = "\n".join(str(r[0]) for r in small[0]["Summary"].iter_rows(max_col=1, values_only=True) if r[0])
    assert "Most pledges: A with 18" in text
    assert "Most productive region: A at 1.50" in text
    assert "Best location vs target: Alpha at 150%" in text
    assert "1 Scale or repeat" in text and "1 Keep and monitor" in text


def test_region_sheet_values_and_formats(small):
    ws = small[0]["By Region"]
    a = next(r for r in table(ws) if r["Region"] == "A")
    assert a["Pledges"] == 18 and a["Doors knocked"] == 400
    assert a["Conversation rate"] == pytest.approx(0.225) and a["Pledge rate"] == pytest.approx(0.2)
    assert a["Pledges per hour"] == pytest.approx(1.5) and a["Avg monthly donation (EUR)"] == pytest.approx(15.0)
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=headers.index("Conversation rate") + 1).number_format == "0.0%"
    assert ws.cell(row=2, column=headers.index("Pledges per hour") + 1).number_format == "0.00"
    assert [r["Region"] for r in table(ws)] == ["A", "B"]  # most pledges first


def test_locations_sheet_attainment_and_classification(small):
    rows = {r["Location ID"]: r for r in table(small[0]["Locations"])}
    assert rows["L1"]["Target attainment"] == pytest.approx(1.5) and rows["L1"]["Planning tier"] == "Scale or repeat"
    assert rows["L2"]["Target attainment"] == pytest.approx(0.8) and rows["L2"]["Planning tier"] == "Keep and monitor"
    assert rows["L2"]["PPH gap vs target"] == pytest.approx(-0.2)  # 0.8 - 1.0
    assert "80%" in rows["L2"]["Why"]


def test_data_issues_sheet_lists_the_bad_row(small):
    issues = table(small[0]["Data Issues"])
    assert len(issues) == 1
    i = issues[0]
    assert (i["Dataset"], i["Source row"], i["Rule"], i["Severity"]) == ("activity", 5, "conversations_gt_doors", "error")
    assert i["Column"] == "conversations" and "20" in i["What is wrong"]


@pytest.mark.parametrize("sheet", ["By Region", "Locations", "Data Issues"])
def test_table_sheets_have_frozen_header_and_filter(small, sheet):
    ws = small[0][sheet]
    assert ws.freeze_panes in ("A2", "B2")  # header row always frozen
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A1:")
    assert ws["A1"].font.bold and ws["A1"].fill.start_color.rgb.endswith("1F3A5F")


@pytest.mark.parametrize("sheet", ["By Region", "Locations", "Data Issues"])
def test_every_column_is_wide_enough_to_read(small, sheet):
    ws = small[0][sheet]
    for cell in ws[1]:
        width = ws.column_dimensions[cell.column_letter].width
        assert width >= 11, f"{sheet}!{cell.value} too narrow ({width})"
    # long text columns get real room
    headers = {c.value: c.column_letter for c in ws[1]}
    for long_col in ("What is wrong", "Why"):
        if long_col in headers:
            assert ws.column_dimensions[headers[long_col]].width >= 50


def conditional_rules(ws):
    return {str(cf.sqref): [r.type for r in cf.rules] for cf in ws.conditional_formatting}


def test_conditional_formatting_is_applied(small):
    wb, _, _ = small
    region = conditional_rules(wb["By Region"])
    assert any("colorScale" in v for v in region.values())
    assert any("dataBar" in v for v in region.values())
    loc = conditional_rules(wb["Locations"])
    assert any("colorScale" in v for v in loc.values())  # attainment
    priority = next(cf for cf in wb["Locations"].conditional_formatting if cf.rules[0].type == "cellIs")
    labels = {r.formula[0] for r in priority.rules}
    assert {'"Scale or repeat"', '"Keep and monitor"', '"Review before repeating"', '"Insufficient data"'} <= labels
    sev = next(iter(wb["Data Issues"].conditional_formatting))
    assert {r.formula[0] for r in sev.rules} == {'"error"', '"warning"'}


def test_methods_sheet_documents_definitions_and_settings(small):
    text = "\n".join(" ".join(str(c) for c in r if c) for r in small[0]["Methods"].iter_rows(values_only=True))
    assert "conversations ÷ doors knocked" in text and "achieved PPH ÷ target PPH" in text
    assert "Minimum hours" in text and "Scale or repeat" in text and "never average row-level rates" in text


def test_missing_values_are_blank_cells_not_zero_or_text(tmp_path):
    (tmp_path / "s.csv").write_text("region,pledges,active_hours,doors_knocked,conversations\nA,0,3,0,0\n", encoding="utf-8")
    cfg = config_from_dict({"inputs": {"activity": "s.csv"}}, tmp_path)
    ws = load_workbook(write_report(run(cfg), cfg, tmp_path / "r.xlsx"))["By Region"]
    row = table(ws)[0]
    assert row["Conversation rate"] is None and row["Pledge rate"] is None  # 0 / 0 is unknown, not 0%
    assert row["Pledges per hour"] == 0  # 0 / 3 really is zero


def test_report_without_location_input_still_builds(tmp_path):
    (tmp_path / "s.csv").write_text("region,pledges,active_hours\nA,4,2\n", encoding="utf-8")
    cfg = config_from_dict({"inputs": {"activity": "s.csv"}}, tmp_path)
    wb = load_workbook(write_report(run(cfg), cfg, tmp_path / "r.xlsx"))
    assert "No location data" in wb["Locations"]["A1"].value
    assert wb["Data Issues"]["A2"].value == "No data issues found."
    assert summary_values(wb["Summary"])["Pledges per hour"] == pytest.approx(2.0)


def test_full_sample_report(tmp_path):
    cfg = load_config(ROOT / "config" / "sample.yaml")
    result = run(cfg)
    wb = load_workbook(write_report(result, cfg, tmp_path / "sample.xlsx"))
    assert wb["Data Issues"].max_row == 1 + 14  # header + the 14 documented errors
    assert wb["By Region"].max_row == 1 + 5
    assert wb["Locations"].max_row == 1 + 30
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and v.lower() in ("nan", "none", "nat") for v in row)
    prio = [r["Planning tier"] for r in table(wb["Locations"])]
    assert prio == list(result.by_location["planning_priority"])


# ---- text from the user's file must never become a spreadsheet formula ---------------------------------

HOSTILE = '=HYPERLINK("http://example.invalid","open me")'


def test_text_starting_with_an_equals_sign_stays_plain_text_on_every_sheet(tmp_path):
    (tmp_path / "s.csv").write_text(
        "date,region,activity_type,fundraiser_id,location_id,pledges,active_hours\n"
        f'2024-01-08,"{HOSTILE.replace(chr(34), chr(34) * 2)}",Stand,"=1+1",L1,5,2\n'
        '2024-01-09,B,"=cmd|\' /c calc\'!A0",F2,L1,x,2\n', encoding="utf-8")
    (tmp_path / "l.csv").write_text(f'location_id,location_label,target_pph\nL1,"=1+2",1.0\n', encoding="utf-8")
    cfg = config_from_dict({"inputs": {"activity": "s.csv", "locations": "l.csv"},
                            "classification": {"min_hours": 1}}, tmp_path)
    wb = load_workbook(write_report(run(cfg), cfg, tmp_path / "r.xlsx"))
    formulas = [(ws.title, c.coordinate) for ws in wb for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert formulas == []
    # and the text itself is kept exactly as it was in the file
    assert HOSTILE in [r["Region"] for r in table(wb["By Region"])]
    assert "=1+1" in [r["Value"] for r in table(wb["By Dimension"])]
    assert "=1+2" in [r["Location"] for r in table(wb["Locations"])]
    assert any(i["Value found"] == "x" for i in table(wb["Data Issues"]))
