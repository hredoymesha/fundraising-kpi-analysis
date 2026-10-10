"""report / profile / init-config, the column-matching helper, and the shipped config files."""
import os
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from f2f_kpi import load_config, run
from f2f_kpi.cli import main
from f2f_kpi.profile import columns_yaml, suggest_mapping

ROOT = Path(__file__).resolve().parents[1]
SHIFTS = str(ROOT / "data" / "sample" / "shifts.csv")
LOCATIONS = str(ROOT / "data" / "sample" / "locations.csv")
DEFAULT_MAPPING = ROOT / "config" / "default_mapping.yaml"
TEMPLATE = ROOT / "src" / "f2f_kpi" / "templates" / "config.example.yaml"

TEAM_CSV = ("Datum,Gebiet,Sign-ups,Hours worked,Doors,Convos,Monthly EUR,Notes\n"
            "2024-01-09,Nord,5,4.5,100,30,75,x\n2024-01-10,Nord,7,5.5,150,45,105,y\n")


def mapping(columns, dataset="activity"):
    return {s.column: s.field for s in suggest_mapping(columns, dataset)[0]}


# ---- column matching --------------------------------------------------------------------------------

def test_headers_in_another_style_are_matched():
    got = mapping(["Datum", "Gebiet", "Sign-ups", "Hours worked", "Doors", "Convos", "Monthly EUR", "Notes"])
    assert got == {"Datum": "date", "Gebiet": "region", "Sign-ups": "pledges", "Hours worked": "active_hours",
                   "Doors": "doors_knocked", "Convos": "conversations", "Monthly EUR": "total_monthly_value_eur"}


def test_matching_ignores_case_accents_and_punctuation():
    assert mapping(["REGION", "Active_Hours", "Pledges!!"]) == {
        "REGION": "region", "Active_Hours": "active_hours", "Pledges!!": "pledges"}


def test_a_column_or_field_is_used_only_once():
    suggestions, _, unused = suggest_mapping(["Hours", "Hours worked", "Region", "Pledges"])
    assert [s.field for s in suggestions].count("active_hours") == 1
    assert len(unused) == 1 and unused[0] in ("Hours", "Hours worked")  # the other one is left for a human to look at


def test_unrelated_columns_are_left_alone_and_missing_required_fields_reported():
    suggestions, missing, unused = suggest_mapping(["Colour", "Region", "Pledges"])
    assert unused == ["Colour"]
    assert missing == ["active_hours"]  # region and pledges found, hours is not
    assert {s.field for s in suggestions} == {"region", "pledges"}


def test_other_datasets_use_their_own_fields():
    got = mapping(["Venue ID", "Venue", "Target", "Week", "Pledges", "Hours"], "location_weekly")
    assert got["Venue ID"] == "location_id" and got["Target"] == "target_pph" and got["Week"] == "period"


def test_confidence_labels():
    by = {s.column: s for s in suggest_mapping(["Pledges", "Hrs wrked", "Region"])[0]}
    assert by["Pledges"].confidence == "exact"
    assert by["Hrs wrked"].confidence in ("likely", "check this one")


def test_columns_yaml_lists_only_headers_that_differ_and_quotes_safely():
    s = suggest_mapping(["region", "Sign-ups", 'Hours "worked"'])[0]
    block = yaml.safe_load(columns_yaml(s))["columns"]
    assert block == {"activity": {"pledges": "Sign-ups", "active_hours": 'Hours "worked"'}}  # 'region' needs no entry
    assert columns_yaml(suggest_mapping(["region", "pledges", "active_hours"])[0]) == "columns: {}"


# ---- profile ----------------------------------------------------------------------------------------

@pytest.fixture
def team_csv(tmp_path):
    p = tmp_path / "team.csv"
    p.write_text(TEAM_CSV, encoding="utf-8")
    return p


def test_profile_prints_suggestions_examples_and_yaml(team_csv, capsys):
    assert main(["profile", str(team_csv)]) == 0
    out = capsys.readouterr().out
    assert "Gebiet" in out and "region" in out and "e.g. Nord" in out
    assert "Columns not used: Notes" in out
    assert 'pledges: "Sign-ups"' in out


def test_profile_exit_code_is_1_when_a_required_field_has_no_match(tmp_path, capsys):
    p = tmp_path / "odd.csv"
    p.write_text("Colour,Size\nred,1\n", encoding="utf-8")
    assert main(["profile", str(p)]) == 1
    assert "Required fields with no matching column" in capsys.readouterr().out


def test_profile_reads_excel_sheet_and_header_row(tmp_path, capsys):
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily"
    ws.append(["Weekly export"])
    ws.append(["Area", "Sign-ups", "Hrs"])
    ws.append(["N", 3, 2])
    wb.save(tmp_path / "b.xlsx")
    assert main(["profile", str(tmp_path / "b.xlsx"), "--sheet", "Daily", "--header-row", "2"]) == 0
    assert "Area" in capsys.readouterr().out


def test_profile_missing_file_exits_2(tmp_path, capsys):
    assert main(["profile", str(tmp_path / "nope.csv")]) == 2
    assert "nope.csv" in capsys.readouterr().err


# ---- init-config ------------------------------------------------------------------------------------

def test_init_config_without_a_file_writes_the_template(tmp_path, capsys):
    out = tmp_path / "cfg" / "mine.yaml"
    assert main(["init-config", "-o", str(out)]) == 0
    assert out.read_text(encoding="utf-8") == TEMPLATE.read_text(encoding="utf-8")


def test_init_config_will_not_overwrite_unless_forced(tmp_path, capsys):
    out = tmp_path / "c.yaml"
    out.write_text("keep me", encoding="utf-8")
    assert main(["init-config", "-o", str(out)]) == 2
    assert out.read_text(encoding="utf-8") == "keep me" and "--force" in capsys.readouterr().err
    assert main(["init-config", "-o", str(out), "--force"]) == 0
    assert out.read_text(encoding="utf-8") != "keep me"


def test_init_config_with_a_data_file_produces_a_config_that_runs(team_csv, tmp_path, capsys):
    out = tmp_path / "team.yaml"
    assert main(["init-config", str(team_csv), "-o", str(out)]) == 0
    text = out.read_text(encoding="utf-8")
    assert "activity: team.csv" in text and 'pledges: "Sign-ups"' in text
    assert "# activity:\n  #   region" not in text  # the template's placeholder example is gone
    res = run(load_config(out))
    assert res.issues.empty
    # hours 4.5 + 5.5 = 10 ; pledges 5 + 7 = 12
    assert res.overall.loc[0, "pledges_per_hour"] == pytest.approx(1.2)
    assert res.overall.loc[0, "total_monthly_value_eur"] == 180


def test_init_config_tells_you_what_it_could_not_match(tmp_path, capsys):
    p = tmp_path / "odd.csv"
    p.write_text("Region,Pledges\nA,1\n", encoding="utf-8")
    assert main(["init-config", str(p), "-o", str(tmp_path / "c.yaml")]) == 0
    assert "Could not match: active_hours" in capsys.readouterr().out


# ---- report -----------------------------------------------------------------------------------------

def test_report_with_the_documented_command(tmp_path, capsys):
    out = tmp_path / "r.xlsx"
    code = main(["report", SHIFTS, "--config", str(DEFAULT_MAPPING), "-o", str(out)])
    assert code == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "By Region", "By Dimension", "Locations", "Trends", "Data Issues", "Methods"]
    assert "activity: 1557 rows read" in capsys.readouterr().out
    assert len(wb["Trends"]._charts) == 3
    assert "No location data" in wb["Locations"]["A1"].value  # no location list given


def test_report_works_without_any_config(tmp_path):
    assert main(["report", SHIFTS, "-o", str(tmp_path / "r.xlsx")]) == 0


def test_report_with_a_location_list_fills_the_locations_sheet(tmp_path):
    out = tmp_path / "r.xlsx"
    assert main(["report", SHIFTS, "--locations", LOCATIONS, "-o", str(out)]) == 0
    ws = load_workbook(out)["Locations"]
    assert ws.max_row == 1 + 30 and ws["A1"].value == "Location ID"


def test_report_input_replaces_the_path_in_the_config_but_keeps_its_other_settings(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily"
    ws.append(["Weekly export"])
    ws.append([])
    ws.append(["Area", "Sign-ups", "Hours"])  # header on row 3
    ws.append(["N", 6, 2])
    wb.save(tmp_path / "team.xlsx")
    (tmp_path / "cfg.yaml").write_text(
        "inputs:\n  activity: {path: somewhere_else.xlsx, sheet: Daily, header_row: 3}\n"
        "columns:\n  activity:\n    region: Area\n    pledges: Sign-ups\n    active_hours: Hours\n", encoding="utf-8")
    out = tmp_path / "r.xlsx"
    assert main(["report", str(tmp_path / "team.xlsx"), "-c", str(tmp_path / "cfg.yaml"), "-o", str(out)]) == 0
    s = {r[0]: r[1] for r in load_workbook(out)["Summary"].iter_rows(max_col=2, values_only=True) if r[0]}
    assert s["Pledges per hour"] == pytest.approx(3.0)  # 6 / 2


def test_report_strict_fails_on_the_sample_but_still_writes_the_file(tmp_path, capsys):
    out = tmp_path / "r.xlsx"
    assert main(["report", SHIFTS, "-o", str(out), "--strict"]) == 1
    assert out.exists() and "Strict mode" in capsys.readouterr().err


def test_report_with_a_missing_input_exits_2(tmp_path, capsys):
    assert main(["report", str(tmp_path / "gone.csv"), "-o", str(tmp_path / "r.xlsx")]) == 2
    assert "gone.csv" in capsys.readouterr().err


def test_run_and_validate_still_need_a_config(capsys):
    for cmd in ("run", "validate"):
        with pytest.raises(SystemExit) as exc:
            main([cmd])
        assert exc.value.code == 2


# ---- the shipped config files -----------------------------------------------------------------------

def test_default_mapping_is_the_packaged_template():
    assert DEFAULT_MAPPING.read_text(encoding="utf-8") == TEMPLATE.read_text(encoding="utf-8")


def test_template_is_a_valid_config():
    cfg = load_config(DEFAULT_MAPPING)
    assert cfg.report.trend_period == "week" and cfg.validation.max_hours_per_row == 14
    assert [t.label for t in cfg.classification.tiers] == ["Scale or repeat", "Keep and monitor", "Review before repeating"]


def test_every_commented_example_in_the_template_is_valid_when_switched_on(tmp_path):
    """Un-comment the optional lines and make sure they load: examples should not rot."""
    lines = []
    for line in TEMPLATE.read_text(encoding="utf-8").splitlines():
        stripped = line.lstrip()
        if stripped.startswith("# ") and ":" in stripped and not stripped.startswith("# Not sure") and \
                not stripped.startswith(("# Copy", "# f2f", "# Paths", "# For ", "# Tell", "# The left", "# Number", "# Translate",
                                         "# Decide", "# Stricter", "#   f2f")):
            line = line.replace("# ", "", 1)
        lines.append(line)
    text = "\n".join(lines).replace("columns: {}", "columns:")
    (tmp_path / "all.yaml").write_text(text, encoding="utf-8")
    cfg = load_config(tmp_path / "all.yaml")
    assert cfg.report.dimensions == ["activity_type", "fundraiser_id"]
    assert cfg.classification.tiers_by_type["Fair"][0].min == 1.2


def test_module_entry_point_runs_the_documented_command(tmp_path):
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
    out = tmp_path / "r.xlsx"
    done = subprocess.run([sys.executable, "-m", "f2f_kpi", "report", SHIFTS, "--config", str(DEFAULT_MAPPING),
                           "--output", str(out)], capture_output=True, text=True, env=env, cwd=ROOT)
    assert done.returncode == 0, done.stderr
    assert out.exists()
