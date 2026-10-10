"""By Dimension and Trends: engine tables, then the workbook sheets and native charts.

Hand-calculated data (one row per shift, Mondays so each row is its own week unless noted):
  2024-01-08 (W02): doors 100, conv 40, pledges 8,  value 120, hours 4      -> PPH 2.0
  2024-01-15 (W03): doors 200, conv 50, pledges 10, value 160, hours 10     -> PPH 1.0
  2024-01-22 (W04): nothing (a quiet week)
  2024-01-29 (W05): doors 100, conv 20, pledges 6,  value 90,  hours 3      -> PPH 2.0
  2024-02-05 (W06): doors 100, conv 30, pledges 12, value 180, hours 6      -> PPH 2.0
"""
import math

import pandas as pd
import pytest
from openpyxl import load_workbook

from f2f_kpi import compute_trends, config_from_dict, run, write_report
from f2f_kpi.config import ConfigError

CSV = (
    "date,region,activity_type,fundraiser_id,doors_knocked,conversations,pledges,total_monthly_value_eur,active_hours\n"
    "2024-01-08,N,Stand,F1,100,40,8,120,4\n"
    "2024-01-15,N,Stand,F1,200,50,10,160,10\n"
    "2024-01-29,S,Door to Door,F2,100,20,6,90,3\n"
    "2024-02-05,S,Door to Door,F2,100,30,12,180,6\n"
)


@pytest.fixture
def cfg(tmp_path):
    (tmp_path / "s.csv").write_text(CSV, encoding="utf-8")
    return lambda extra=None: config_from_dict({"inputs": {"activity": "s.csv"}, **(extra or {})}, tmp_path)


# ---- engine: trends ---------------------------------------------------------------------------------

def test_weekly_trends_by_hand(cfg):
    t = run(cfg()).trends.set_index("period")
    assert list(t.index) == ["2024-W02", "2024-W03", "2024-W04", "2024-W05", "2024-W06"]  # the quiet week is a row
    assert t.loc["2024-W02", "pledges_per_hour"] == pytest.approx(2.0)  # 8 / 4
    assert t.loc["2024-W03", "pledges_per_hour"] == pytest.approx(1.0)  # 10 / 10
    assert t.loc["2024-W03", "conversation_rate"] == pytest.approx(0.25)  # 50 / 200
    assert t.loc["2024-W03", "pledge_rate"] == pytest.approx(0.2)  # 10 / 50
    assert t.loc["2024-W03", "shifts"] == 1 and t.loc["2024-W04", "shifts"] == 0
    assert t.loc["2024-W02", "period_start"] == pd.Timestamp("2024-01-08")


def test_change_against_previous_period(cfg):
    t = run(cfg()).trends.set_index("period")
    assert math.isnan(t.loc["2024-W02", "pledges_change_pct"])  # nothing before the first week
    assert t.loc["2024-W03", "pledges_change_pct"] == pytest.approx((10 - 8) / 8)  # +25%
    assert t.loc["2024-W03", "pph_change_pct"] == pytest.approx((1.0 - 2.0) / 2.0)  # -50%


def test_quiet_week_is_blank_and_the_week_after_has_no_change(cfg):
    t = run(cfg()).trends.set_index("period")
    assert math.isnan(t.loc["2024-W04", "pledges_per_hour"])  # no hours worked: unknown, not 0
    assert t.loc["2024-W04", "pledges"] == 0 or math.isnan(t.loc["2024-W04", "pledges"])
    assert math.isnan(t.loc["2024-W05", "pph_change_pct"])  # previous week had no data to compare with


def test_rolling_four_period_pph_is_a_ratio_of_sums(cfg):
    t = run(cfg()).trends.set_index("period")
    assert t["pph_rolling_4"].iloc[:3].isna().all()  # needs four periods first
    # W02..W05: pledges 8 + 10 + 0 + 6 = 24, hours 4 + 10 + 0 + 3 = 17  (not the mean of 2.0, 1.0, 2.0 = 1.67)
    assert t.loc["2024-W05", "pph_rolling_4"] == pytest.approx(24 / 17)
    # W03..W06: pledges 10 + 0 + 6 + 12 = 28, hours 10 + 0 + 3 + 6 = 19
    assert t.loc["2024-W06", "pph_rolling_4"] == pytest.approx(28 / 19)


def test_monthly_trends_group_by_calendar_month(cfg):
    t = run(cfg({"report": {"trend_period": "month"}})).trends.set_index("period")
    assert list(t.index) == ["2024-01", "2024-02"]
    # January: pledges 8 + 10 + 6 = 24, hours 4 + 10 + 3 = 17 ; February: 12 / 6
    assert t.loc["2024-01", "pledges"] == 24 and t.loc["2024-01", "pledges_per_hour"] == pytest.approx(24 / 17)
    assert t.loc["2024-02", "pledges_per_hour"] == pytest.approx(2.0)
    assert t.loc["2024-02", "pledges_change_pct"] == pytest.approx((12 - 24) / 24)  # -50%


def test_no_date_column_means_no_trends(tmp_path):
    (tmp_path / "s.csv").write_text("region,pledges,active_hours\nA,4,2\n", encoding="utf-8")
    assert run(config_from_dict({"inputs": {"activity": "s.csv"}}, tmp_path)).trends is None
    assert compute_trends(pd.DataFrame({"pledges": [1]})).empty


def test_bad_trend_settings_are_rejected():
    with pytest.raises(ConfigError, match="trend_period"):
        config_from_dict({"inputs": {"activity": "x.csv"}, "report": {"trend_period": "day"}})
    with pytest.raises(ValueError):
        compute_trends(pd.DataFrame({"date": pd.to_datetime(["2024-01-01"])}), "day")


# ---- engine: dimensions -----------------------------------------------------------------------------

def test_dimension_kpis_by_hand(cfg):
    d = run(cfg()).by_dimension.set_index(["dimension", "value"])
    # Stand: doors 300, conv 90, pledges 18, value 280, hours 14
    stand = d.loc[("activity_type", "Stand")]
    assert stand["pledges"] == 18 and stand["shifts"] == 2
    assert stand["conversation_rate"] == pytest.approx(90 / 300)
    assert stand["pledges_per_hour"] == pytest.approx(18 / 14)
    assert stand["avg_monthly_donation_eur"] == pytest.approx(280 / 18)
    # Door to Door: doors 200, conv 50, pledges 18, value 270, hours 9
    d2d = d.loc[("activity_type", "Door to Door")]
    assert d2d["pledge_rate"] == pytest.approx(18 / 50) and d2d["pledges_per_hour"] == pytest.approx(2.0)
    assert d.loc[("fundraiser_id", "F2"), "pledges"] == 18


def test_dimensions_can_be_chosen_in_the_config(cfg):
    only = run(cfg({"report": {"dimensions": ["fundraiser_id"]}})).by_dimension
    assert set(only["dimension"]) == {"fundraiser_id"}
    with pytest.raises(ConfigError, match="cannot be used"):
        config_from_dict({"inputs": {"activity": "x.csv"}, "report": {"dimensions": ["region"]}})


def test_missing_dimension_columns_are_skipped(tmp_path):
    (tmp_path / "s.csv").write_text("region,pledges,active_hours\nA,4,2\n", encoding="utf-8")
    assert run(config_from_dict({"inputs": {"activity": "s.csv"}}, tmp_path)).by_dimension is None


# ---- workbook ---------------------------------------------------------------------------------------

def rows_of(ws):
    r = list(ws.iter_rows(values_only=True))
    return [dict(zip(r[0], x)) for x in r[1:] if x[0] is not None]


@pytest.fixture
def book(cfg, tmp_path):
    c = cfg()
    return load_workbook(write_report(run(c), c, tmp_path / "r.xlsx"))


def test_by_dimension_sheet(book):
    ws = book["By Dimension"]
    rows = rows_of(ws)
    assert {(r["Breakdown"], r["Value"]) for r in rows[:2]} == {("Activity type", "Stand"), ("Activity type", "Door to Door")}
    stand = next(r for r in rows if r["Value"] == "Stand")
    assert stand["Pledges"] == 18 and stand["Pledges per hour"] == pytest.approx(18 / 14)
    assert {r["Breakdown"] for r in rows} == {"Activity type", "Fundraiser"}
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    assert any(rule.type == "colorScale" for cf in ws.conditional_formatting for rule in cf.rules)


def test_each_breakdown_is_shaded_on_its_own(book):
    ranges = [str(cf.sqref) for cf in book["By Dimension"].conditional_formatting]
    # 2 activity types in rows 2-3, 2 fundraisers in rows 4-5: no scale spans both blocks
    assert ranges and all(r.split(":")[0][1:] in ("2", "4") for r in ranges)


def test_trends_sheet_values_and_formats(book):
    ws = book["Trends"]
    rows = rows_of(ws)
    assert [r["Period"] for r in rows] == ["2024-W02", "2024-W03", "2024-W04", "2024-W05", "2024-W06"]
    w3 = rows[1]
    assert w3["Pledges per hour"] == pytest.approx(1.0) and w3["Pledges vs previous"] == pytest.approx(0.25)
    assert w3["PPH vs previous"] == pytest.approx(-0.5)
    assert rows[2]["Pledges per hour"] is None  # the quiet week stays blank
    headers = {c.value: c for c in ws[1]}
    assert ws.cell(row=3, column=headers["Pledges vs previous"].column).number_format == "+0.0%;-0.0%;0.0%"
    assert ws.freeze_panes == "B2"
    rules = [r.operator for cf in ws.conditional_formatting for r in cf.rules if r.type == "cellIs"]
    assert {"greaterThan", "lessThan"} <= set(rules)


def test_trends_sheet_has_three_native_charts_pointing_at_the_table(book):
    ws = book["Trends"]
    assert len(ws._charts) == 3
    kinds = sorted(type(c).__name__ for c in ws._charts)
    assert kinds == ["BarChart", "LineChart", "LineChart"]
    for chart in ws._charts:
        for s in chart.series:
            assert ws.title in s.val.numRef.f and "$" in s.val.numRef.f  # refers to cells, not pasted numbers
        assert chart.anchor._from.row >= 5 + 3  # below the 5-row table, never covering it


def test_charts_survive_a_reload(book, tmp_path):
    out = tmp_path / "again.xlsx"
    book.save(out)
    assert len(load_workbook(out)["Trends"]._charts) == 3


def test_sheets_degrade_to_a_plain_note_when_there_is_nothing_to_show(tmp_path):
    (tmp_path / "s.csv").write_text("region,pledges,active_hours\nA,4,2\n", encoding="utf-8")
    c = config_from_dict({"inputs": {"activity": "s.csv"}}, tmp_path)
    wb = load_workbook(write_report(run(c), c, tmp_path / "r.xlsx"))
    assert "No trends" in wb["Trends"]["A1"].value and not wb["Trends"]._charts
    assert "No breakdown columns" in wb["By Dimension"]["A1"].value


def test_sample_data_trends_and_dimensions(tmp_path):
    from pathlib import Path
    from f2f_kpi import load_config
    c = load_config(Path(__file__).resolve().parents[1] / "config" / "sample.yaml")
    res = run(c)
    assert len(res.trends) == 12  # twelve weeks, no gaps
    assert res.by_dimension.groupby("dimension").size().to_dict() == {"activity_type": 3, "fundraiser_id": 40}
    wb = load_workbook(write_report(res, c, tmp_path / "s.xlsx"))
    assert wb["Trends"].max_row == 1 + 12 and len(wb["Trends"]._charts) == 3
    assert wb["By Dimension"].max_row == 1 + 43
