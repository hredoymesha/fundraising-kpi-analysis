"""The web app, driven through Streamlit's own test harness: real widgets, real uploads, real engine.

Every test also checks that nothing crashed: no exception element and no 'Something unexpected' error box
(the app turns surprise exceptions into that box, so looking only for exceptions would miss them).
"""
import io
import socket
from datetime import date
from pathlib import Path

import pytest

pytest.importorskip("streamlit")
pytest.importorskip("plotly")

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from f2f_kpi import analyse
from f2f_kpi.io import read_bytes

import app_logic as logic

ROOT = Path(__file__).resolve().parents[1]
APP = str(ROOT / "app" / "streamlit_app.py")

# Hand-calculated (see test_filters_and_bytes.py for the same shape):
#   N: 2024-01-08 doors 100 conv 30 pledges 6 value 90 hours 4 ; 2024-01-09 doors 300 conv 60 pledges 12 value 180 hours 8
#   S: 2024-01-15 doors 200 conv 40 pledges 4 value 80 hours 5
TWO_REGIONS = (b"date,region,location_id,doors_knocked,conversations,pledges,total_monthly_value_eur,active_hours\n"
               b"2024-01-08,N,L1,100,30,6,90,4\n2024-01-09,N,L1,300,60,12,180,8\n2024-01-15,S,L2,200,40,4,80,5\n")
LOCATIONS = b"location_id,location_label,location_type,target_pph\nL1,Alpha,Stand,1.0\nL2,Beta,Fair,1.0\n"
# a team that names its columns differently: pledges 5 + 7 = 12, hours 4.5 + 5.5 = 10, doors 250, conv 75, value 180
TEAM = (b"Datum,Gebiet,Sign-ups,Hours worked,Doors,Convos,Monthly EUR,Notes\n"
        b"2024-01-09,Nord,5,4.5,100,30,75,x\n2024-01-10,Nord,7,5.5,150,45,105,y\n")


def start() -> AppTest:
    return AppTest.from_file(APP, default_timeout=90).run()


def upload(at: AppTest, key: str, name: str, data: bytes) -> AppTest:
    boxes = at.sidebar.get("file_uploader")
    if not boxes or not hasattr(boxes[0], "upload"):
        pytest.skip("this Streamlit's test harness cannot upload files (the dev extra asks for streamlit>=1.65)")
    box = next(u for u in boxes if u.key == key)
    box.upload(name, data)
    return at.run()


def with_file(data: bytes, name: str = "export.csv", locations: bytes | None = None) -> AppTest:
    at = upload(start(), "up_activity", name, data)
    if locations is not None:
        at = upload(at, "up_locations", "locations.csv", locations)
    return at


def sample() -> AppTest:
    at = start()
    next(b for b in at.sidebar.button if b.label == "Try it with sample data").click()
    return at.run()


def healthy(at: AppTest) -> None:
    assert not at.exception, [e.value for e in at.exception]
    assert not [e for e in at.error if "Something unexpected" in e.value], [e.value for e in at.error]


def metrics(at: AppTest) -> dict:
    return {m.label: m.value for m in at.metric}


def select(at: AppTest, label: str):
    return next(s for s in at.selectbox if s.label == label)


# ---- first visit ------------------------------------------------------------------------------------

def test_first_visit_explains_what_to_do_and_offers_the_sample():
    at = start()
    healthy(at)
    assert at.title[0].value == "F2F fundraising KPIs"
    assert any("Upload" in m.value for m in at.markdown)
    assert [b.label for b in at.sidebar.button] == ["Try it with sample data"]
    assert not at.metric and not at.get("download_button")


# ---- the sample data end to end ---------------------------------------------------------------------

def test_sample_data_shows_cards_tabs_and_download():
    at = sample()
    healthy(at)
    assert at.success[0].value.startswith("Read 1557 rows from your activity file and used 1549")
    m = metrics(at)
    assert list(m)[:5] == ["Conversation rate", "Pledge rate", "Pledges per hour", "Average monthly gift",
                           "Target attainment"]  # the cards come first; tier counts and issue counters follow
    cfg = logic.default_config()
    expected = analyse(cfg, {"activity": read_bytes((ROOT / "data/sample/shifts.csv").read_bytes(), "s.csv"),
                             "locations": read_bytes((ROOT / "data/sample/locations.csv").read_bytes(), "l.csv")})
    cards = {c.label: c.value for c in logic.kpi_cards(expected)}
    for label, value in cards.items():
        assert m[label] == value  # the page shows what the engine says, formatted
    assert m["Pledges per hour"] == "1.84" and m["Target attainment"] == "112%"
    assert [t.label for t in at.tabs] == ["Overview", "Locations", "Data issues (12)"]
    assert [b.label for b in at.get("download_button")] == ["Download Excel report"]


def test_sample_charts_are_drawn():
    at = sample()
    healthy(at)
    assert len(at.get("plotly_chart")) == 4  # pph by region, rates by region, two trend charts


def test_locations_tab_ranks_every_location():
    at = sample()
    table = at.tabs[1].dataframe[0].value
    assert len(table) == 30 and list(table["Rank"]) == list(range(1, 31))
    assert table["Target attainment (%)"].is_monotonic_decreasing
    assert {"Location", "Planning tier", "Why"} <= set(table.columns)


def test_data_issues_tab_lists_source_rows():
    at = sample()
    table = at.tabs[2].dataframe[0].value
    assert len(table) == 12
    assert list(table["Severity"]).count("error") == 8 and list(table["Severity"]).count("warning") == 4
    assert table["Row in file"].notna().all()
    assert "conversations_gt_doors" in set(table["Rule"])


def test_stopping_the_sample_returns_to_the_start():
    at = sample()
    next(b for b in at.sidebar.button if b.label == "Stop using the sample").click()
    at.run()
    at.run()  # some Streamlit versions keep showing the old elements for one run after st.rerun()
    healthy(at)
    assert not at.metric and [b.label for b in at.sidebar.button] == ["Try it with sample data"]


# ---- uploading and column mapping -------------------------------------------------------------------

def test_upload_with_unfamiliar_headers_is_matched_and_calculated():
    at = with_file(TEAM, "team.csv")
    healthy(at)
    assert select(at, "Region (required)").value == "Gebiet"
    assert select(at, "Pledges (sign-ups) (required)").value == "Sign-ups"
    assert select(at, "Active hours (required)").value == "Hours worked"
    assert select(at, "Total monthly value (EUR)").value == "Monthly EUR"
    m = metrics(at)
    assert m["Conversation rate"] == "30.0%"  # 75 / 250
    assert m["Pledge rate"] == "16.0%"  # 12 / 75
    assert m["Pledges per hour"] == "1.20"  # 12 / 10
    assert m["Average monthly gift"] == "€15.00"  # 180 / 12
    assert m["Target attainment"] == "—"  # no location list yet


def test_missing_required_column_asks_for_it_then_works_once_chosen():
    at = with_file(b"Gebiet,Sign-ups,Zeit\nNord,6,2\nNord,4,2\n")
    healthy(at)
    assert any("Still needed" in w.value and "Active hours" in w.value for w in at.warning)
    assert not at.metric and not at.get("download_button")
    select(at, "Active hours (required)").set_value("Zeit")
    at.run()
    healthy(at)
    assert metrics(at)["Pledges per hour"] == "2.50"  # 10 pledges / 4 hours
    assert not any("Still needed" in w.value for w in at.warning)


def test_one_column_cannot_fill_two_fields():
    at = with_file(b"region,pledges,active_hours\nA,4,2\n")
    select(at, "Active hours (required)").set_value("pledges")
    at.run()
    healthy(at)
    assert any("only one field" in w.value and "'pledges'" in w.value for w in at.warning)
    assert not at.metric


def test_an_optional_field_can_be_switched_off():
    at = with_file(TWO_REGIONS)
    assert select(at, "Date").value == "date"
    select(at, "Date").set_value(logic.NOT_MAPPED)
    at.run()
    healthy(at)
    assert not [c for c in at.get("plotly_chart") if "by week" in str(c.proto.spec)]
    assert any("Trends need a date column" in i.value for i in at.info)


# ---- bad files --------------------------------------------------------------------------------------

@pytest.mark.parametrize("name,data,fragment", [
    ("broken.csv", bytes([0xff, 0xfe, 0x00, 0x81]) * 60, "not readable as utf-8"),
    ("empty.csv", b"  \n", "is empty"),
    ("fake.xlsx", b"this is not a workbook", "not a valid Excel workbook"),
])
def test_unreadable_files_show_a_plain_error_and_stop(name, data, fragment):
    at = with_file(data, name)
    healthy(at)
    assert any(fragment in e.value for e in at.error), [e.value for e in at.error]
    assert not at.metric and not at.get("download_button")


def test_a_file_with_headings_but_no_rows_says_so():
    at = with_file(b"region,pledges,active_hours\n")
    healthy(at)
    assert any("no data rows" in e.value for e in at.error)
    assert not at.metric


def test_when_every_row_is_unusable_the_problems_are_shown_instead_of_empty_charts():
    at = with_file(b"region,pledges,active_hours\nA,lots,2\nB,many,3\n")
    healthy(at)
    assert any("None of the 2 rows could be used" in e.value for e in at.error)
    assert "Pledges per hour" not in metrics(at) and metrics(at)["Errors"] == "2"
    assert not at.get("download_button") and not at.get("plotly_chart")
    table = at.dataframe[0].value
    assert list(table["Rule"]) == ["non_numeric_value", "non_numeric_value"] and list(table["Row in file"]) == [2, 3]


def test_a_broken_location_list_is_skipped_but_the_activity_results_stay():
    at = upload(with_file(TWO_REGIONS), "up_locations", "locations.csv", bytes([0xff, 0xfe, 0x81]) * 40)
    healthy(at)
    assert any("Location list" in e.value for e in at.error)
    assert any("without the location list" in w.value for w in at.warning)
    assert metrics(at)["Pledges per hour"] == "1.29"  # 22 / 17
    assert metrics(at)["Target attainment"] == "—"


def test_location_list_with_a_missing_required_column_is_left_out_with_a_hint():
    at = with_file(TWO_REGIONS, locations=b"location_id,name\nL1,Alpha\n")
    healthy(at)
    assert any("Still needed" in w.value and "Target pledges per hour" in w.value for w in at.warning)
    assert any("left out until this is fixed" in c.value for c in at.caption)
    assert metrics(at)["Pledges per hour"] == "1.29"


# ---- number / date format, settings file, Excel -----------------------------------------------------

def test_comma_decimals_need_the_decimal_mark_setting():
    data = b"region;pledges;active_hours\nA;9;4,5\n"
    at = with_file(data, "eu.csv")
    healthy(at)
    assert any("None of the 1 rows could be used" in e.value for e in at.error)  # 4,5 is not a number yet
    at.sidebar.radio[0].set_value("Comma (4,5)")
    at.run()
    healthy(at)
    assert metrics(at)["Pledges per hour"] == "2.00"  # 9 / 4.5


def test_a_settings_file_changes_tiers_and_defaults():
    at = with_file(TWO_REGIONS, locations=LOCATIONS)
    before = at.tabs[1].dataframe[0].value
    assert set(before["Planning tier"]) == {"Insufficient data"}  # 12 and 5 hours: under the default 20-hour minimum
    yaml_box = next(u for u in at.sidebar.get("file_uploader") if u.key == "up_config")
    yaml_box.upload("rules.yaml", b"classification:\n  min_hours: 1\n")
    at.run()
    healthy(at)
    after = at.tabs[1].dataframe[0].value
    assert list(after["Planning tier"]) == ["Scale or repeat", "Keep and monitor"]  # 150% and 80%
    assert list(after["Target attainment (%)"]) == pytest.approx([150.0, 80.0])


def test_an_invalid_settings_file_is_reported_and_defaults_are_used():
    at = with_file(TWO_REGIONS)
    next(u for u in at.sidebar.get("file_uploader") if u.key == "up_config").upload("bad.yaml", b"colums: {}")
    at.run()
    healthy(at)
    assert any("Could not use that settings file" in e.value and "unknown top-level" in e.value for e in at.sidebar.error)
    assert metrics(at)["Pledges per hour"] == "1.29"


def _workbook() -> bytes:
    wb = Workbook()
    wb.active.title = "Cover"
    wb.active.append(["nothing to see"])
    daily = wb.create_sheet("Daily")
    daily.append(["Weekly export"])
    daily.append(["region", "pledges", "active_hours"])
    daily.append(["A", 9, 4.5])
    daily.append(["A", 3, 1.5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_upload_with_sheet_and_header_row_choices():
    at = with_file(_workbook(), "book.xlsx")
    healthy(at)  # sheet 'Cover' and row 1: no usable data yet
    assert not at.metric
    at.sidebar.selectbox(key="sheet_activity").set_value("Daily")
    at.sidebar.number_input(key="hdr_activity").set_value(2)
    at.run()
    healthy(at)
    assert metrics(at)["Pledges per hour"] == "2.00"  # 12 pledges / 6 hours


# ---- filters ----------------------------------------------------------------------------------------

def test_region_filter_recalculates_the_cards():
    at = with_file(TWO_REGIONS)
    assert metrics(at)["Pledges per hour"] == "1.29"  # 22 / 17
    at.sidebar.multiselect[0].set_value(["N"])
    at.run()
    healthy(at)
    m = metrics(at)
    assert m["Conversation rate"] == "22.5%" and m["Pledge rate"] == "20.0%"  # 90/400 ; 18/90
    assert m["Pledges per hour"] == "1.50" and m["Average monthly gift"] == "€15.00"  # 18/12 ; 270/18
    assert any("Showing 2 shifts" in c.value and "1 of 2 regions" in c.value for c in at.caption)


def test_date_filter_includes_both_end_dates():
    at = with_file(TWO_REGIONS)
    at.sidebar.date_input[0].set_value((date(2024, 1, 9), date(2024, 1, 15)))
    at.run()
    healthy(at)
    assert metrics(at)["Pledges per hour"] == "1.23"  # (12 + 4) / (8 + 5)
    assert any("Showing 2 shifts" in c.value and "09 Jan 2024 to 15 Jan 2024" in c.value for c in at.caption)


def test_filters_that_match_nothing_get_a_clear_message():
    at = with_file(TWO_REGIONS)
    at.sidebar.multiselect[0].set_value([])
    at.run()
    healthy(at)
    assert any("No rows match these filters" in w.value for w in at.warning)
    assert not at.metric


def test_a_half_chosen_date_range_does_not_break_the_page():
    at = with_file(TWO_REGIONS)
    at.sidebar.date_input[0].set_value((date(2024, 1, 9),))
    at.run()
    healthy(at)
    assert any("Pick an end date" in c.value for c in at.sidebar.caption)


# ---- the Excel download -----------------------------------------------------------------------------

def test_download_serves_a_real_workbook_that_matches_the_data():
    at = with_file(TWO_REGIONS, locations=LOCATIONS)
    assert len(at.get("download_button")) == 1
    data = at.session_state["analysis"]["report"]
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "By Region", "By Dimension", "Locations", "Trends", "Data Issues", "Methods"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(max_col=2, values_only=True) if r[0]}
    assert summary["Pledges"] == 22 and summary["Pledges per hour"] == pytest.approx(22 / 17)
    assert wb["Locations"].max_row == 3  # header + 2 locations
    assert len(wb["Trends"]._charts) == 3


def test_the_report_covers_everything_even_when_filters_are_set():
    at = with_file(TWO_REGIONS)
    before = at.session_state["analysis"]["report"]
    at.sidebar.multiselect[0].set_value(["N"])
    at.run()
    after = at.session_state["analysis"]["report"]
    assert after is before  # not rebuilt, not narrowed
    summary = {r[0]: r[1] for r in load_workbook(io.BytesIO(after))["Summary"].iter_rows(max_col=2, values_only=True) if r[0]}
    assert summary["Pledges"] == 22
    assert any("covers every usable row" in c.value for c in at.sidebar.caption)


def test_the_download_includes_the_data_issues_found():
    at = sample()
    wb = load_workbook(io.BytesIO(at.session_state["analysis"]["report"]))
    assert wb["Data Issues"].max_row == 1 + 12


# ---- privacy ----------------------------------------------------------------------------------------

def test_the_app_makes_no_network_connections(monkeypatch):
    outside = []
    real_connect = socket.socket.connect
    real_create = socket.create_connection

    def is_local(address) -> bool:
        host = address[0] if isinstance(address, tuple) else str(address)
        return host in ("127.0.0.1", "::1", "localhost", "")

    def guarded_connect(self, address):
        if not is_local(address):
            outside.append(address)
            raise AssertionError("the app tried to connect to something outside this computer")
        return real_connect(self, address)

    def guarded_create(address, *a, **k):
        if not is_local(address):
            outside.append(address)
            raise AssertionError("the app tried to connect to something outside this computer")
        return real_create(address, *a, **k)

    monkeypatch.setattr(socket.socket, "connect", guarded_connect)
    monkeypatch.setattr(socket, "create_connection", guarded_create)
    at = with_file(TEAM, "team.csv")
    healthy(at)
    assert metrics(at)["Pledges per hour"] == "1.20" and outside == []


def test_the_app_writes_no_files_except_the_spreadsheet_librarys_temporary_ones(monkeypatch):
    """openpyxl streams each sheet through a temp file while building a report; it deletes them again."""
    import glob
    import os
    import tempfile

    written = []
    real_open = open

    def watching_open(file, mode="r", *a, **k):
        if any(c in mode for c in "wax+") and "__pycache__" not in str(file):
            written.append(os.path.basename(str(file)))
        return real_open(file, mode, *a, **k)

    pattern = os.path.join(tempfile.gettempdir(), "openpyxl.*")
    before = set(glob.glob(pattern))
    monkeypatch.setattr("builtins.open", watching_open)
    at = with_file(TEAM, "team.csv")
    at.run()
    assert metrics(at)["Pledges per hour"] == "1.20"
    assert [name for name in written if not name.startswith("openpyxl.")] == []
    assert set(glob.glob(pattern)) == before  # nothing left behind in the temp folder


def test_app_code_has_no_networking_logging_or_disk_writes():
    banned = ["import requests", "urllib", "http.client", "import socket", "import logging", "tempfile", "subprocess",
              "os.system", "print(", ".write_text(", ".write_bytes(", "open("]
    for path in [ROOT / "app" / "streamlit_app.py", ROOT / "app" / "app_logic.py", ROOT / "app" / "figures.py"]:
        text = path.read_text(encoding="utf-8")
        for token in banned:
            assert token not in text, f"{path.name} contains {token!r}"


def test_streamlit_config_keeps_the_app_local_and_quiet():
    text = (ROOT / ".streamlit" / "config.toml").read_text(encoding="utf-8")
    assert "gatherUsageStats = false" in text
    assert 'address = "localhost"' in text
    assert 'showErrorDetails = "type"' in text


def test_a_crash_never_echoes_the_user_s_data(monkeypatch):
    import f2f_kpi

    def explode(*a, **k):
        raise RuntimeError("SecretTownName 12345 leaked")

    monkeypatch.setattr(f2f_kpi, "analyse", explode)
    at = with_file(TEAM, "team.csv")
    assert not at.exception
    shown = " ".join(e.value for e in at.error)
    assert "Something unexpected went wrong (RuntimeError)" in shown
    assert "SecretTownName" not in shown and "12345" not in shown


# ---- reachability --------------------------------------------------------------------------------------

def _pretend_server_address(monkeypatch, value):
    import streamlit as st
    real = st.get_option
    monkeypatch.setattr(st, "get_option", lambda key: value if key == "server.address" else real(key))


def test_no_network_warning_when_the_server_is_limited_to_this_computer(monkeypatch):
    _pretend_server_address(monkeypatch, "localhost")
    at = start()
    healthy(at)
    assert not [w for w in at.warning if "Other computers" in w.value]


def test_the_app_warns_when_it_is_reachable_from_the_network(monkeypatch):
    _pretend_server_address(monkeypatch, None)  # what Streamlit reports when started outside the repo root
    for at in (start(), with_file(TWO_REGIONS)):  # on the welcome page and once results are showing
        healthy(at)
        warning = next(w for w in at.warning if "Other computers" in w.value)
        assert "--server.address localhost" in warning.value


def test_this_repository_s_own_settings_limit_the_server_to_localhost():
    import streamlit as st
    # pytest runs from the repository root, where .streamlit/config.toml applies
    assert st.get_option("server.address") == "localhost"
    assert st.get_option("browser.gatherUsageStats") is False


# ---- text from the user's files must show as text -----------------------------------------------------

IMAGE_HEADER = "![x](https://evil.example/p.png)"


def test_a_crafted_column_header_is_not_rendered_as_an_image_in_the_warning():
    at = with_file(('region,"' + IMAGE_HEADER + '",active_hours' + chr(10) + "A,4,2" + chr(10)).encode())
    select(at, "Pledges (sign-ups) (required)").set_value(IMAGE_HEADER)
    at.run()
    select(at, "Conversations").set_value(IMAGE_HEADER)
    at.run()
    healthy(at)
    warning = next(w.value for w in at.warning if "only one field" in w.value)
    assert IMAGE_HEADER not in warning and "![x](" not in warning  # the raw Markdown never reaches the page
    assert warning.replace(chr(92), "").count(IMAGE_HEADER) == 1  # but the reader still sees the header's text


def test_a_crafted_file_name_is_not_rendered_as_markdown_in_labels_or_errors():
    name = "![probe](p.png).csv"
    ok = with_file(TWO_REGIONS, name)
    healthy(ok)
    assert not any(name in e.label for e in ok.expander)  # the raw Markdown never reaches the page
    assert any(name in e.label.replace(chr(92), "") for e in ok.expander)  # but the reader still sees the name
    broken = with_file(bytes([0xff, 0xfe, 0x00, 0x81]) * 60, name)
    healthy(broken)
    assert all(name not in e.value for e in broken.error)
    assert any("not readable as utf-8" in e.value for e in broken.error)
