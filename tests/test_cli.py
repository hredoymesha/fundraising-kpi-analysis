import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from f2f_kpi.cli import main

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = str(ROOT / "config" / "sample.yaml")


@pytest.fixture
def clean_config(tmp_path):
    (tmp_path / "s.csv").write_text("region,pledges,active_hours\nA,4,2\nB,3,3\n", encoding="utf-8")
    (tmp_path / "c.yaml").write_text("inputs:\n  activity: s.csv\n", encoding="utf-8")
    return str(tmp_path / "c.yaml")


def test_run_writes_the_report_and_prints_a_summary(tmp_path, capsys):
    out = tmp_path / "nested" / "report.xlsx"
    assert main(["run", "--config", SAMPLE, "--output", str(out)]) == 0
    assert out.exists() and load_workbook(out).sheetnames[0] == "Summary"
    text = capsys.readouterr().out
    assert "Report written to" in text
    assert "activity: 1557 rows read, 1545 used" in text
    assert "12 error(s)" in text and "2 warning(s)" in text


def test_run_default_output_name(tmp_path, clean_config, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert main(["run", "-c", clean_config]) == 0
    assert (tmp_path / "f2f_kpi_report.xlsx").exists()


def test_strict_fails_on_issues_but_still_writes_the_report(tmp_path, capsys):
    out = tmp_path / "r.xlsx"
    assert main(["run", "-c", SAMPLE, "-o", str(out), "--strict"]) == 1
    assert out.exists()
    assert "Strict mode" in capsys.readouterr().err


def test_strict_passes_on_clean_data(tmp_path, clean_config):
    assert main(["run", "-c", clean_config, "-o", str(tmp_path / "r.xlsx"), "--strict"]) == 0


def test_non_strict_run_succeeds_even_with_data_errors(tmp_path):
    assert main(["run", "-c", SAMPLE, "-o", str(tmp_path / "r.xlsx")]) == 0  # bad rows are excluded and reported


def test_validate_lists_issues_with_rows_and_fails_on_errors(capsys):
    assert main(["validate", "-c", SAMPLE]) == 1
    text = capsys.readouterr().out
    assert "conversations_gt_doors" in text and "activity row 210" in text
    assert "WARNING" in text and "ERROR" in text


def test_validate_clean_data_exits_zero_and_writes_no_report(tmp_path, clean_config, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert main(["validate", "-c", clean_config]) == 0
    assert "No issues found." in capsys.readouterr().out
    assert not list(tmp_path.glob("*.xlsx"))


def test_validate_warnings_only_pass_unless_strict(tmp_path, capsys):
    (tmp_path / "s.csv").write_text("region,pledges,active_hours\nA,4,2\na,3,3\nA,1,1\n", encoding="utf-8")  # 'a' vs 'A'
    (tmp_path / "c.yaml").write_text("inputs:\n  activity: s.csv\n", encoding="utf-8")
    assert main(["validate", "-c", str(tmp_path / "c.yaml")]) == 0
    assert main(["validate", "-c", str(tmp_path / "c.yaml"), "--strict"]) == 1


def test_missing_config_is_a_clean_error_not_a_traceback(tmp_path, capsys):
    assert main(["run", "-c", str(tmp_path / "nope.yaml"), "-o", str(tmp_path / "r.xlsx")]) == 2
    err = capsys.readouterr().err
    assert err.startswith("error:") and "Traceback" not in err


def test_bad_config_and_missing_data_file_exit_2(tmp_path, capsys):
    (tmp_path / "bad.yaml").write_text("inputs:\n  activity: a.csv\nwat: 1\n", encoding="utf-8")
    assert main(["validate", "-c", str(tmp_path / "bad.yaml")]) == 2
    assert "unknown top-level" in capsys.readouterr().err
    (tmp_path / "ok.yaml").write_text("inputs:\n  activity: missing.csv\n", encoding="utf-8")
    assert main(["validate", "-c", str(tmp_path / "ok.yaml")]) == 2
    assert "missing.csv" in capsys.readouterr().err


def test_no_command_is_a_usage_error(capsys):
    with pytest.raises(SystemExit) as exc:
        main([])
    assert exc.value.code == 2


def test_output_that_cannot_be_written_exits_2(tmp_path, clean_config, capsys):
    blocker = tmp_path / "blocker"
    blocker.mkdir()  # a directory where the file should go
    assert main(["run", "-c", clean_config, "-o", str(blocker)]) == 2
    assert "error:" in capsys.readouterr().err


def test_python_dash_m_entry_point(tmp_path):
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
    out = tmp_path / "r.xlsx"
    done = subprocess.run([sys.executable, "-m", "f2f_kpi", "run", "-c", SAMPLE, "-o", str(out)],
                          capture_output=True, text=True, env=env)
    assert done.returncode == 0, done.stderr
    assert out.exists() and "Report written to" in done.stdout
    version = subprocess.run([sys.executable, "-m", "f2f_kpi", "--version"], capture_output=True, text=True, env=env)
    assert version.stdout.strip() == "f2f-kpi 0.1.0"


# ---- found in the final review -------------------------------------------------------------------------

def test_report_with_no_usable_rows_says_why_and_does_not_claim_success(tmp_path, capsys):
    src = tmp_path / "other_headers.csv"
    src.write_text("Gebiet,Zusagen,Stunden\nNord,5,4\n", encoding="utf-8")
    out = tmp_path / "r.xlsx"
    assert main(["report", str(src), "-o", str(out)]) == 1
    err = capsys.readouterr().err
    assert "Nothing could be used" in err and "No column found for the required field" in err
    assert out.exists()  # the report is still written: its Data Issues sheet explains what went wrong


def test_report_that_only_loses_some_rows_still_succeeds(tmp_path):
    src = tmp_path / "mixed.csv"
    src.write_text("region,pledges,active_hours\nA,5,2\nA,oops,2\n", encoding="utf-8")
    assert main(["report", str(src), "-o", str(tmp_path / "r.xlsx")]) == 0


def test_a_missing_config_file_is_named_in_plain_words(tmp_path, capsys):
    assert main(["validate", "-c", str(tmp_path / "nope.yaml")]) == 2
    err = capsys.readouterr().err
    assert err.startswith("error: config file not found:") and "Errno" not in err
